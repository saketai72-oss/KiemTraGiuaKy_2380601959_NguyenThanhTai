import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def create_tagged_exam_docx(output_path):
    doc = docx.Document()
    
    # Title & Duration tags
    p = doc.add_paragraph()
    p.add_run("[EXAM_TITLE: ETS TOEIC Reading Test 09]").bold = True
    
    p = doc.add_paragraph()
    p.add_run("[DURATION: 75]").bold = True
    doc.add_paragraph()

    # ==========================================
    # PART 5: QUESTIONS 101 - 130
    # ==========================================
    doc.add_paragraph("[PART: 5]")
    
    part5_data = [
        (101, "Every batch of sauce at Generita's Bistro is processed meticulously by ------- expert chefs.",
         [("A", "they"), ("B", "their"), ("C", "them"), ("D", "themselves")], "B",
         "Cần tính từ sở hữu 'their' đứng trước cụm danh từ 'expert chefs'."),
        (102, "Computerization of medical records ------- increases a physician's ability to diagnose and treat patients.",
         [("A", "great"), ("B", "greatly"), ("C", "greatness"), ("D", "greatest")], "B",
         "Cần trạng từ 'greatly' để bổ nghĩa cho động từ 'increases'."),
        (103, "At Rojelle's Fine Dining, we use the freshest ------- available to make our salads.",
         [("A", "applications"), ("B", "subjects"), ("C", "ingredients"), ("D", "factors")], "C",
         "'ingredients' có nghĩa là nguyên liệu, thành phần chế biến món ăn."),
        (104, "Professor Benguigui will present his paper ------- the natural history conference tomorrow.",
         [("A", "by"), ("B", "at"), ("C", "of"), ("D", "on")], "B",
         "Giới từ 'at' đi với sự kiện/hội nghị: 'at the conference'."),
        (105, "The chefs are still waiting for a ------- date for the new convection oven that was ordered two weeks ago.",
         [("A", "shipper"), ("B", "ships"), ("C", "shipments"), ("D", "shipping")], "D",
         "Cụm danh từ 'shipping date' có nghĩa là ngày giao/vận chuyển hàng."),
        (106, "Mr. Kang works ------- with our internal team members as well as various regional sales representatives.",
         [("A", "mildly"), ("B", "nearly"), ("C", "closely"), ("D", "narrowly")], "C",
         "Cụm từ 'work closely with' có nghĩa là làm việc chặt chẽ, hợp tác mật thiết với ai."),
        (107, "Fashion designer Hye-Ja Pak knows ------- to update her line in response to changing tastes.",
         [("A", "and"), ("B", "when"), ("C", "need"), ("D", "for")], "B",
         "Cấu trúc từ để hỏi + to-V: 'knows when to update' (biết khi nào nên cập nhật)."),
        (108, "The samba class was so well ------- that the Yulara Dance School decided to make the course a permanent offering.",
         [("A", "attended"), ("B", "educated"), ("C", "gathered"), ("D", "protected")], "A",
         "Cụm tính từ 'well attended' mang nghĩa có đông đảo học viên tham gia."),
        (109, "The factory will be situated away ------- the city's residential area to reduce complaints about noise and emissions.",
         [("A", "from"), ("B", "about"), ("C", "with"), ("D", "out")], "A",
         "Cụm giới từ 'away from' mang nghĩa nằm cách xa khỏi khu vực nào."),
        (110, "Ms. Kuramoto selected the most ------- mailing option available.",
         [("A", "economical"), ("B", "economy"), ("C", "economize"), ("D", "economized")], "A",
         "Cần tính từ 'economical' (tiết kiệm) sau 'the most' để bổ nghĩa cho cụm 'mailing option'."),
        (111, "Companies without information technology specialists can ------- on Vyber Software Advisers for assistance with online services.",
         [("A", "reliable"), ("B", "reliably"), ("C", "rely"), ("D", "relying")], "C",
         "Sau trợ động từ 'can' cần động từ nguyên mẫu không 'to': 'rely on' (trông cậy vào)."),
        (112, "Next year, our team will have a new task, ------- is to review design portfolios.",
         [("A", "although"), ("B", "which"), ("C", "after"), ("D", "because")], "B",
         "Đại từ quan hệ 'which' thay thế cho danh từ 'a new task' ở mệnh đề phía trước."),
        (113, "Cook the meat for 30 minutes to ensure ------- readiness to be eaten.",
         [("A", "both"), ("B", "this"), ("C", "its"), ("D", "that")], "C",
         "Tính từ sở hữu 'its' thay thế cho 'the meat' đứng trước danh từ 'readiness'."),
        (114, "Amoxitron's research team will hire ------- interns to assist with laboratory duties.",
         [("A", "given"), ("B", "several"), ("C", "whole"), ("D", "natural")], "B",
         "'Several' (một vài) đi với danh từ đếm được số nhiều 'interns'."),
        (115, "------- the new acai juice blend has proved so popular, we should move quickly to increase our production volume.",
         [("A", "If"), ("B", "Whether"), ("C", "Since"), ("D", "Unless")], "C",
         "Liên từ 'Since' mang nghĩa 'bởi vì' chỉ nguyên nhân phù hợp với ngữ cảnh."),
        (116, "New emissions standards have forced Rider Auto to modify the process of engine -------.",
         [("A", "construction"), ("B", "constructed"), ("C", "constructive"), ("D", "construct")], "A",
         "Sau giới từ 'of' cần cụm danh từ 'engine construction' (sự chế tạo động cơ)."),
        (117, "Maki Kayano's book offers techniques for ------- business tasks with speed and precision.",
         [("A", "executing"), ("B", "equipping"), ("C", "returning"), ("D", "involving")], "A",
         "'executing' mang nghĩa thực thi/thực hiện các nhiệm vụ kinh doanh."),
        (118, "To control costs, updated credit card readers will be installed in branch stores -------.",
         [("A", "gradual"), ("B", "gradually"), ("C", "more gradual"), ("D", "most gradual")], "B",
         "Trạng từ 'gradually' (dần dần, từng bước) bổ nghĩa cho cụm động từ 'will be installed'."),
        (119, "In this quarter, the Montel Beverage Company is expecting sales ------- £160,000 and £180,000.",
         [("A", "without"), ("B", "among"), ("C", "throughout"), ("D", "between")], "D",
         "Cấu trúc liên từ 'between ... and ...' biểu thị khoảng dao động giữa hai con số."),
        (120, "Because Mount Akoyola is so challenging for climbers, ------- have reached its peak.",
         [("A", "any"), ("B", "either"), ("C", "other"), ("D", "few")], "D",
         "Đại từ 'few' mang nghĩa phủ định (rất ít người), phù hợp với lý do ngọn núi rất hiểm trở."),
        (121, "King Street Bridge will be closed in the month of September ------- repair work.",
         [("A", "because of"), ("B", "so that"), ("C", "as if"), ("D", "rather than")], "A",
         "'because of' đi với cụm danh từ 'repair work' để chỉ lý do đóng cầu."),
        (122, "Ms. Taniguchi's supervisor commended her for negotiating ------- with Furuyama Corporation.",
         [("A", "effective"), ("B", "effecting"), ("C", "effected"), ("D", "effectively")], "D",
         "Trạng từ 'effectively' bổ nghĩa cho danh động từ 'negotiating'."),
        (123, "Staff members may reserve the conference room ------- they need it.",
         [("A", "somewhere"), ("B", "whatever"), ("C", "everything"), ("D", "anytime")], "D",
         "Liên từ chỉ thời gian 'anytime' mang nghĩa bất cứ khi nào họ cần."),
        (124, "The public relations director must have a high level of ------- in English and Spanish.",
         [("A", "proficiency"), ("B", "advancement"), ("C", "routine"), ("D", "strength")], "A",
         "Danh từ 'proficiency' đi với giới từ 'in' mang nghĩa sự thông thạo ngôn ngữ."),
        (125, "Patterson Products ------- seeks innovative ways of meeting changing consumer demand.",
         [("A", "dually"), ("B", "favorably"), ("C", "continually"), ("D", "generically")], "C",
         "Trạng từ 'continually' mang nghĩa liên tục, không ngừng tìm kiếm giải pháp sáng tạo."),
        (126, "Mr. Volante is working at home tomorrow so he ------- the technical report without any distractions.",
         [("A", "can finish"), ("B", "would finish"), ("C", "finished"), ("D", "has been finishing")], "A",
         "Mệnh đề chỉ mục đích diễn tả khả năng có thể hoàn thành: 'so he can finish'."),
        (127, "All of the billing procedures remain the same, ------- new payment codes need to be inserted into the invoice documents.",
         [("A", "in order that"), ("B", "during"), ("C", "across from"), ("D", "except that")], "D",
         "Liên từ 'except that' mang nghĩa ngoại trừ việc cần bổ sung mã thanh toán mới."),
        (128, "The lightweight design of the new sedan is ------- because it can cause the vehicle to slide on icy roads.",
         [("A", "possible"), ("B", "mechanical"), ("C", "questionable"), ("D", "multiple")], "C",
         "Tính từ 'questionable' mang nghĩa đáng nghi ngại, cần xem xét lại do gây nguy cơ trượt tuyết."),
        (129, "When the bank president retired, the common ------- was that the vice president would take over.",
         [("A", "assume"), ("B", "assumption"), ("C", "assuming"), ("D", "assumable")], "B",
         "Sau tính từ 'common' cần danh từ 'assumption' làm chủ ngữ của mệnh đề."),
        (130, "West Bengali Airlines ------- fees for oversized items that are still within weight limitations.",
         [("A", "invites"), ("B", "cooperates"), ("C", "transports"), ("D", "waives")], "D",
         "Động từ 'waives' mang nghĩa miễn trừ khoản phí (waive fees).")
    ]
    
    for q_num, text, options, key, hint in part5_data:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: TRUE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # ==========================================
    # PART 6: QUESTIONS 131 - 146
    # ==========================================
    doc.add_paragraph("[PART: 6]")
    doc.add_paragraph()

    # Passage 1: Questions 131-134
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 131-134 refer to the following e-mail.")
    doc.add_paragraph("To: pmendoza@factmail.co\nFrom: recruiting@analystsassoc.org\nDate: May 2\nSubject: Joining ASA\n\nDear Mr. Mendoza,\nThank you for expressing your interest in our organization during the recent [131]. It was a pleasure meeting you at the ASA booth during the Weber Information Systems Convention. As you may recall from our conversation, we discussed how ASA membership [132] your career through networking opportunities as well as the industry insights offered in our monthly newsletter. The normal fee for members is just $120 a year; however, we are currently offering a new-member discount. [133].\nIf you are still interested in joining, please reply to [134] with your mailing address. I will then forward you an application packet.\n\nSincerely,\nAshlee Loren, President\nAssociation of Systems Analysts")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part6_p1 = [
        (131, "131.", [("A", "election"), ("B", "broadcast"), ("C", "conference"), ("D", "performance")], "C", "Đoạn văn đề cập đến 'Weber Information Systems Convention' nên từ phù hợp là 'conference' (hội nghị)."),
        (132, "132.", [("A", "can benefit"), ("B", "is benefiting"), ("C", "has benefited"), ("D", "will have benefited")], "A", "Diễn tả khả năng lợi ích mà việc tham gia mang lại cho sự nghiệp: 'can benefit'."),
        (133, "133.", [("A", "Discounts on rental cars are included in the annual membership."), ("B", "Our jobs board is quite comprehensive."), ("C", "We are the first organization of our kind."), ("D", "This month you can join for just half the usual rate.")], "D", "Câu trước nhắc đến 'new-member discount', nên câu tiếp theo mô tả mức giảm giá: chỉ một nửa giá thông thường."),
        (134, "134.", [("A", "it"), ("B", "me"), ("C", "them"), ("D", "anyone")], "B", "Người gửi dùng 'I will then forward you' nên câu trước đề nghị gửi phản hồi cho tôi: 'reply to me'.")
    ]
    for q_num, text, options, key, hint in part6_p1:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Passage 2: Questions 135-138
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 135-138 refer to the following e-mail.")
    doc.add_paragraph("To: Sam Heinz\nFrom: Northways Professional Development\nSent: April 20\nSubject: Workshop 4/28-4/30\n\nDear Workshop Participants,\nWe look forward to seeing you at the digital storytelling workshop. On the first day, when you enter the Albin College campus, attendants will [135] you to lot 43 and the Toteman Building. We will begin each day in conference room 9. Coffee, tea, snacks, and fruit [136] in the mornings. Lunch will be sandwiches and salads from Black Horse Restaurant. [137].\nIn the afternoons, we will be working in the computer lab. We suggest that you assemble some images that you would like to use for your project. It will be [138] for you to have them saved on a digital storage device beforehand. If you have any questions, please e-mail us.\n\nVery Best,\nGina Kapuski")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part6_p2 = [
        (135, "135.", [("A", "offer"), ("B", "direct"), ("C", "pass"), ("D", "instruct")], "B", "'direct you to' có nghĩa là chỉ dẫn, hướng dẫn ai đó đến vị trí cụ thể."),
        (136, "136.", [("A", "will be provided"), ("B", "were provided"), ("C", "providers"), ("D", "are providing")], "A", "Thì tương lai bị động 'will be provided' diễn tả sự việc sẽ được chuẩn bị trong buổi hội thảo sắp tới."),
        (137, "137.", [("A", "Some prior experience working with digital files is assumed."), ("B", "Please do not bring any copyrighted material to the workshop."), ("C", "Please let us know if you have any dietary restrictions."), ("D", "There is a one-time parking charge of fifteen dollars.")], "C", "Câu trước nói về thực đơn bữa trưa, câu nối tiếp hợp lý nhất là nhắc báo trước nếu có kiêng ăn (dietary restrictions)."),
        (138, "138.", [("A", "useful"), ("B", "surprised"), ("C", "difficult"), ("D", "amazing")], "A", "'useful' mang nghĩa hữu ích khi lưu trữ trước các tài liệu/hình ảnh vào thiết bị nhớ.")
    ]
    for q_num, text, options, key, hint in part6_p2:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Passage 3: Questions 139-142
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 139-142 refer to the following article.")
    doc.add_paragraph("Venley Foods Responsive to Changing Consumers\n\nWho cares where the tomatoes in your salad actually came from? [139], an increasing number of people do, according to a study conducted by the Consumer Group. In fact, the study shows that many consumers would pay an average of 10 percent more when given the exact source of a fresh food product. [140]. Some grocery stores, such as Venley Foods in Boston, have taken advantage of the trend and used it to implement [141] branding and marketing. \"If we can tell a story about our product,\" says Venley Foods CEO Minji Kim, \"then we've added [142] in the minds of consumers.\"")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part6_p3 = [
        (139, "139.", [("A", "Carefully"), ("B", "Apparently"), ("C", "Formerly"), ("D", "Rarely")], "B", "'Apparently' mang nghĩa rõ ràng/dường như là, dùng để khẳng định đáp lại câu hỏi tu từ mở đầu."),
        (140, "140.", [("A", "Fresh food can be refrigerated for up to two days."), ("B", "Many grocery stores have been extending their hours."), ("C", "Most studies are published in consumer magazines."), ("D", "The number increases to 20 percent in large cities.")], "D", "Bổ sung số liệu thống kê nối tiếp câu trước: tỷ lệ người sẵn sàng trả thêm tăng lên tới 20% ở các thành phố lớn."),
        (141, "141.", [("A", "smarts"), ("B", "smartly"), ("C", "smarter"), ("D", "smartness")], "C", "Cần tính từ so sánh hơn 'smarter' bổ nghĩa cho cụm danh từ 'branding and marketing'."),
        (142, "142.", [("A", "value"), ("B", "time"), ("C", "obstacles"), ("D", "bonus")], "A", "'added value' (gia tăng giá trị) trong nhận thức của người tiêu dùng.")
    ]
    for q_num, text, options, key, hint in part6_p3:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Passage 4: Questions 143-146
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 143-146 refer to the following article.")
    doc.add_paragraph("Nylobe, Inc., Announces Newest Development Project\n\nSEATTLE (May 1)—Researchers at technology firm Nylobe, Inc., are working to develop a sensor capable of detecting corrosion caused by environmental exposure. Corrosion is a major contributor to [143] losses in the aircraft industry each year. \"This will be a major [144] for commercial airline fleets,\" says Mel Laveau, Nylobe's CEO. \"The sensor will decrease both labor and maintenance costs without being too expensive.\"\n\nAccording to Ms. Laveau, the sensor will work by detecting corrosion in its early stages, when the problem can be corrected simply by removing the corroded material. [145]. In the structure of large aircraft, some critical joints can be particularly susceptible to corrosion. [146], the sensor can be used to inspect these areas and then target the most likely areas of concern.")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part6_p4 = [
        (143, "143.", [("A", "financially"), ("B", "financed"), ("C", "financial"), ("D", "finances")], "C", "Cần tính từ 'financial' bổ nghĩa cho danh từ 'losses' (những tổn thất tài chính)."),
        (144, "144.", [("A", "balance"), ("B", "examination"), ("C", "expectation"), ("D", "asset")], "D", "'asset' có nghĩa là một tài sản giá trị/lợi thế lớn cho các đội bay."),
        (145, "145.", [("A", "This will reduce the need for making expensive structural repairs."), ("B", "The parts have all been replaced with higher quality materials."), ("C", "The next stage of the project involves scanning the affected areas."), ("D", "Its style and sleek design made it popular with the public.")], "A", "Giải thích lợi ích của việc phát hiện sớm rỉ sét: giúp giảm thiểu chi phí sửa chữa kết cấu lớn đắt đỏ."),
        (146, "146.", [("A", "Meanwhile"), ("B", "Similarly"), ("C", "Otherwise"), ("D", "Fortunately")], "D", "Trạng từ 'Fortunately' (May mắn thay) chỉ giải pháp tích cực cho vấn đề dễ rỉ sét ở các mối nối.")
    ]
    for q_num, text, options, key, hint in part6_p4:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # ==========================================
    # PART 7: QUESTIONS 147 - 200
    # ==========================================
    doc.add_paragraph("[PART: 7]")
    doc.add_paragraph()

    # Group 147-148: Invoice
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 147-148 refer to the following invoice.\n\nInvoice 3987\nShawqi Office Services, Dubai, UAE\nReport requested in English\n18 October\nJenkins Press\nP.O. Box 2291\nDubai, UAE\n\nService:\nOn 14 October, replaced bulb and repaired paper tray on copier per call received on 12 October. Replaced copy ink in two machines. Performed routine yearly maintenance on five copiers per existing service contract.\n\nLabor cost: AED 330.00\nPaper tray: AED 50.00\nBulb: AED 30.00\nCopy ink: AED 220.00\nTotal: AED 630.00\n\nTotal amount must be received by 31 October.\nThank you for your business!")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part7_g1 = [
        (147, "What is indicated about Jenkins Press?",
         [("A", "It has several offices around the world."), ("B", "Its copiers get checked every year."), ("C", "Its office equipment is outdated."), ("D", "It is a new customer of Shawqi Office Services.")], "B",
         "Trong hóa đơn có ghi: 'routine yearly maintenance on five copiers per existing service contract' -> máy photocopy được kiểm tra bảo dưỡng hàng năm."),
        (148, "When is payment due?",
         [("A", "October 12"), ("B", "October 14"), ("C", "October 18"), ("D", "October 31")], "D",
         "Hạn chót thanh toán ghi rõ: 'Total amount must be received by 31 October.'")
    ]
    for q_num, text, options, key, hint in part7_g1:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Group 149-151: Advertisement
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 149-151 refer to the following advertisement.\n\nHarbour View Apartment—Porthmadog, Wales\n\nThis one-bedroom apartment is perfect for a holiday escape! Located in a quiet area, it boasts a patio with a lovely view of the harbour. Recently renovated, the unit includes an eat-in kitchen with stove, refrigerator, microwave, and coffeepot; bathroom with walk-in shower; and a living room with a large-screen TV. Other amenities of the property include:\n• Short distance to restaurants and shops\n• Five-minute walk to the beach\n• Public gardens and historic sites within a 20-minute drive\n• Heat and electricity included\n• Towels and bed linens provided on-site\n• Daily cleaning service available (extra fee)\n• Wireless Internet access (extra fee)\n\nReserve this lovely gem now! Signing a contract by March 30 will reduce the rental cost by 10 percent. To sign a contract, contact Dylan Barrett at dbarrett@telarentals.co.uk.")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part7_g2 = [
        (149, "Who would the advertisement most likely interest?",
         [("A", "Business travelers"), ("B", "Residents of Porthmadog"), ("C", "Property investors"), ("D", "Short-term vacationers")], "D",
         "Bài quảng cáo giới thiệu căn hộ thích hợp cho chuyến nghỉ mát ngắn hạn ('holiday escape')."),
        (150, "What is indicated about the rental fee?",
         [("A", "It includes tours of historic places."), ("B", "It does not cover all of the apartment's features."), ("C", "It includes vouchers to use at local restaurants."), ("D", "It requires a minimum 30 percent deposit in advance.")], "B",
         "Dịch vụ dọn phòng hàng ngày và wifi có tính thêm phụ phí ('extra fee'), nghĩa là giá thuê không bao gồm tất cả các tiện ích."),
        (151, "Why should an individual contact Mr. Barrett by March 30?",
         [("A", "To get a discount on rent"), ("B", "To schedule apartment renovations"), ("C", "To rent the last available property"), ("D", "To sell the property before the end of the season")], "A",
         "Ký hợp đồng trước ngày 30/3 sẽ được giảm 10% chi phí thuê: 'Signing a contract by March 30 will reduce the rental cost by 10 percent.'")
    ]
    for q_num, text, options, key, hint in part7_g2:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Group 152-153: Text message chain
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 152-153 refer to the following text-message chain.\n\nJason Salter (1:45 P.M.)\nLiz, the meeting is starting in fifteen minutes. Where are you?\n\nLiz Ortiz (1:47 P.M.)\nThe train has been stopped on the tracks for a while. There seems to be some kind of problem. I still hope to make the meeting.\n\nJason Salter (1:50 P.M.)\nOkay. I'll save you a seat.\n\nLiz Ortiz (1:59 P.M.)\nThere's just been an announcement. There's a disabled train up ahead. It'll be awhile.\n\nJason Salter (2:00 P.M.)\nDon't worry. If there are any questions for our department, I'll handle them.\n\nLiz Ortiz (2:01 P.M.)\nThanks. I'll call you later.")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part7_g3 = [
        (152, "What is suggested about Mr. Salter?",
         [("A", "He is leading the meeting."), ("B", "He commutes to work by train."), ("C", "He has many questions for Ms. Ortiz."), ("D", "He works in the same department as Ms. Ortiz.")], "D",
         "Jason Salter nhắn 'If there are any questions for our department, I'll handle them' chứng tỏ cùng thuộc một phòng ban với Liz Ortiz."),
        (153, "At 1:59 P.M., what does Ms. Ortiz imply when she writes, \"It'll be awhile\"?",
         [("A", "The meeting is running late."), ("B", "She is still preparing her notes."), ("C", "She will likely miss the meeting."), ("D", "She has not boarded the train yet.")], "C",
         "Cuộc họp bắt đầu lúc 2:00 P.M., mà đến 1:59 P.M. tàu vẫn đang dừng do tàu phía trước gặp sự cố và 'It'll be awhile', ngụ ý cô ấy sẽ bỏ lỡ cuộc họp.")
    ]
    for q_num, text, options, key, hint in part7_g3:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Group 154-155: Email
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 154-155 refer to the following e-mail.\n\nFrom: tbogosian@sardhasconvention.com\nTo: pradalaily@dmcv.com\nDate: 22 March\nSubject: Your reservation\nAttachment: Prada party\n\nDear Ms. Prada,\nThank you for choosing Sardha's Banquet Hall to host your event on 30 August. Unfortunately, West Hall will be under renovation in August, and we will be unable to accommodate your group in that room. However, your party will fit comfortably in East Hall. This room features floor-to-ceiling windows and a full view of the river. The 15 percent deposit we have received will hold the reservation.\n\nAs indicated when you completed your request through our Web site, your total price of $1,600 includes a full-course meal that includes appetizers, soup or salad, main course, and dessert. Attached to this e-mail please find a copy of the menu presented to guests with options for appetizer and main course.\n\nTomas Bogosian\nGeneral Manager")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part7_g4 = [
        (154, "What is one purpose of the e-mail?",
         [("A", "To cancel a reservation"), ("B", "To apologize for an error"), ("C", "To indicate a room change"), ("D", "To request an extra deposit")], "C",
         "Email thông báo chuyển từ West Hall sang East Hall do phòng cũ đang sửa chữa."),
        (155, "What is suggested about Ms. Prada?",
         [("A", "She has paid her bill in full."), ("B", "She eats at Sardha's regularly."), ("C", "She works in the food industry."), ("D", "She made the reservation online.")], "D",
         "Email nêu rõ: 'when you completed your request through our Web site' -> bà đặt chỗ trực tuyến.")
    ]
    for q_num, text, options, key, hint in part7_g4:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Group 156-157: Memo
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 156-157 refer to the following memo.\n\nMEMO\nTo: All Employees\nFrom: IT Supervisor\nSubject: OS Update\nDate: 27 May\n\nAll company computers in the Melbourne and Victoria offices require an update to the operating system. This update will prepare our computers for the new version of our accounting software, which should arrive in the first week of July. Technicians will install the new operating system beginning on 3 June. We expect the process to be completed around 16 June, providing time to work out any bugs with the operating system before the accounting software is uploaded.\n\nCompleting the entire process will require you to log in so that we can verify the system is working as anticipated. Therefore, if you plan to go on holiday during this period, please inform IT Support immediately of the dates you will be out so that we can plan an alternate date to accommodate your schedule.\n\nIf you have any questions, please contact IT Support staff at extension 48.")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part7_g5 = [
        (156, "When is the installation of the operating system expected to be finished?",
         [("A", "In the first week of June"), ("B", "In the middle of June"), ("C", "At the end of June"), ("D", "In the first week of July")], "B",
         "Dự kiến hoàn thành khoảng ngày 16 tháng 6 ('around 16 June' = 'In the middle of June')."),
        (157, "Who is asked to contact IT Support?",
         [("A", "Employees who do not need the accounting software"), ("B", "Accountants already using the new operating system"), ("C", "Staff who work outside the offices"), ("D", "Employees who are taking time off")], "D",
         "Văn bản ghi: 'if you plan to go on holiday during this period, please inform IT Support immediately' (nghỉ phép = taking time off).")
    ]
    for q_num, text, options, key, hint in part7_g5:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Group 158-160: Letter
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 158-160 refer to the following letter.\n\nClarke-Ellis Construction\n#20 Murphy Industrial Park\nSt. Michael BB23028\n1-246-555-0126\nFull-service commercial contractor serving all of Barbados\n\n4 June\nIda Gutierrez\nDarling Cove Inn\nMango Drive\nFolkestone BB24017\n\nDear Ms. Gutierrez,\nThank you for contacting Clarke-Ellis Construction for your roofing project.\n— [1] —. After inspecting the property, I have confirmed that the inn's main roof is in good condition and requires no repairs at this time. However, the roof of the inn's porch appears to be at least twenty years old and is worn beyond repair. — [2] —. Clarke-Ellis Construction can remove and dispose of the existing porch roof and install a new one. The replacement will be comparable to the quality, style, and colour of the inn's main roof. We will use only commercial-grade leak barriers, insulation, and shingles manufactured by West Indies Weatherproofing, Inc. The estimated total cost, inclusive of labour and materials, is $3,260. — [3] —. Additional fees would apply should you want us to make other improvements, such as painting the porch or replacing porch screens.\n\nPlease call me at the number above to discuss scheduling or any questions you have about the work. — [4] —. I hope to hear from you soon.\n\nSincerely,\nGrayson Clarke, Co-owner, Clarke-Ellis Construction")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part7_g6 = [
        (158, "Why did Mr. Clarke write the letter?",
         [("A", "To provide a work proposal"), ("B", "To ask for a project extension"), ("C", "To request an inspection report"), ("D", "To submit a revised cost estimate")], "A",
         "Nội dung bức thư là bản đề xuất công việc thay thế mái hiên và dự toán chi phí."),
        (159, "What is indicated about West Indies Weatherproofing, Inc.?",
         [("A", "It is installing a new porch at Darling Cove Inn."), ("B", "It provides painting services."), ("C", "It produces roofing materials."), ("D", "It is owned by Clarke-Ellis Construction.")], "C",
         "Thư nêu: 'leak barriers, insulation, and shingles manufactured by West Indies Weatherproofing, Inc.' -> chuyên sản xuất vật liệu lợp mái."),
        (160, "In which of the positions marked [1], [2], [3], and [4] does the following sentence best belong?\n\"It must be replaced.\"",
         [("A", "[1]"), ("B", "[2]"), ("C", "[3]"), ("D", "[4]")], "B",
         "Câu trước vị trí [2] nói mái hiên đã 20 năm và hỏng không thể sửa ('worn beyond repair') -> 'It must be replaced' (Nó cần phải được thay thế). Sau đó là phương án thay thế.")
    ]
    for q_num, text, options, key, hint in part7_g6:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Group 161-164: Advertisement
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 161-164 refer to the following advertisement.\n\nPosition: Assistant Editor\nDate Posted: March 15\n\nDescription:\nGoldhorse Press, an independent book publishing company focusing on North American gardening, has provided gardening advice for home gardeners for over 50 years. We are currently seeking an assistant editor to join our expanding team in Charlotte, North Carolina. We offer an excellent benefits package that includes medical and dental insurance.\n\nResponsibilities:\nAssist editors in the acquisition of titles with mainstream appeal; develop manuscripts, collaborate with outside support, and work closely with two managing editors.\n\nRequirements/Qualifications:\n• Bachelor's degree in a related field\n• Minimum of one year of experience in the publishing industry\n• Ability to pay close attention to detail\n• Comfortable working as a team member\n• Some experience in gardening preferred\n• Immediate availability\n\nE-mail résumé and salary requirements to humanresources@goldhorsepress.com.")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part7_g7 = [
        (161, "What title would most likely be published by Goldhorse Press?",
         [("A", "Growing Your Baking Business"), ("B", "A Tour Guide to North Carolina"), ("C", "Planting Perennial Flowers"), ("D", "The Efficient Executive")], "C",
         "Công ty chuyên xuất bản sách hướng dẫn làm vườn trong hơn 50 năm, nên tựa sách 'Planting Perennial Flowers' là phù hợp nhất."),
        (162, "What is indicated about Goldhorse Press?",
         [("A", "It is hiring two assistant editors."), ("B", "It provides insurance to employees."), ("C", "It publishes trade journals."), ("D", "It is a newly established company.")], "B",
         "Công ty cung cấp chế độ bảo hiểm y tế và nha khoa cho nhân viên ('includes medical and dental insurance')."),
        (163, "What is a requirement for the advertised position?",
         [("A", "A passion for gardening"), ("B", "An aptitude for noticing details"), ("C", "A master's degree in a related field"), ("D", "An ability to work without supervision")], "B",
         "Yêu cầu 'Ability to pay close attention to detail' đồng nghĩa với 'An aptitude for noticing details'."),
        (164, "What are applicants asked to do?",
         [("A", "Submit a job application form"), ("B", "Submit a list of references"), ("C", "Indicate availability"), ("D", "Indicate desired pay")], "D",
         "Ứng viên được yêu cầu gửi 'salary requirements' (mức lương mong muốn, tương đương 'desired pay').")
    ]
    for q_num, text, options, key, hint in part7_g7:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Group 165-168: Chat Session
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 165-168 refer to the following online chat session.\n\nDerek Marshall [8:19 A.M.]\nHi, everyone. I wanted to give an update on our merger with Ridgewood, Inc., and see if you have had any meetings.\n\nMai Chung [8:20 A.M.]\nHave we worked out which of our Derek Boutique locations will stay open?\n\nDerek Marshall [8:21 A.M.]\nYes, Derek Boutique will keep 35 stores open with the full line of clothing. The other 12 will move their inventory to the Ridgewood, Inc., locations.\n\nNikita Tamboli [8:22 A.M.]\nWhen should the moves be scheduled? This month?\n\nDerek Marshall [8:23 A.M.]\nNo, the 12 locations don't need to be vacated until the end of next month. Could you schedule this for five weeks from now?\n\nAnthony Rossi [8:24 A.M.]\nI met over lunch with my managers earlier this week, including the two that I just hired.\n\nNikita Tamboli [8:25 A.M.]\nThat would work.\n\nMai Chung [8:27 A.M.]\nI'll be meeting with the managers on my staff, too, to review the transition plan.\n\nDerek Marshall [8:28 A.M.]\nThis is good work, everyone.\n\nAnthony Rossi [8:29 A.M.]\nWhen will the whole process be complete?\n\nDerek Marshall [8:30 A.M.]\nWithin about 6 months.")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part7_g8 = [
        (165, "Where do the writers most likely work?",
         [("A", "At a marketing firm"), ("B", "At a clothing company"), ("C", "At a real estate agency"), ("D", "At a newspaper publisher")], "B",
         "Họ nhắc đến 'Derek Boutique will keep 35 stores open with the full line of clothing' -> công ty thời trang/may mặc."),
        (166, "How many stores will be closing?",
         [("A", "5"), ("B", "6"), ("C", "12"), ("D", "35")], "C",
         "Có 12 cửa hàng sẽ chuyển kho hàng và bàn giao mặt bằng ('the 12 locations don't need to be vacated...')."),
        (167, "At 8:25 A.M., what does Ms. Tamboli most likely mean when she writes, \"That would work\"?",
         [("A", "The current inventory will be doubled."), ("B", "Ridgewood, Inc., will close in two weeks."), ("C", "Mr. Marshall will meet with the managers."), ("D", "She can schedule the moves in the proposed time frame.")], "D",
         "Derek Marshall hỏi có thể sắp xếp việc di dời sau 5 tuần nữa được không, Tamboli trả lời 'That would work' để xác nhận mốc thời gian này khả thi."),
        (168, "What is suggested about Ms. Chung and Mr. Rossi?",
         [("A", "They supervise other employees."), ("B", "They often shop at Ridgewood, Inc."), ("C", "They recently went to lunch together."), ("D", "They have concerns about the merger.")], "A",
         "Cả Rossi ('my managers') và Chung ('the managers on my staff') đều quản lý cấp dưới trực tiếp.")
    ]
    for q_num, text, options, key, hint in part7_g8:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Group 169-171: Article
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 169-171 refer to the following article.\n\nFrom the Shadows to the Limelight\nBy Calum Ellwood\n\nDr. Esther Nujoma, an agricultural biotechnologist with the Namibia Institute of Applied Sciences, is the author of several books on the practical applications of biotechnology. — [1] —. According to book critic Paige Kinnock of the London Daily Register, \"Dr. Nujoma has increased the public's awareness of the role of biotechnology in daily life through her ability to translate highly complex scientific material into simple language.\"\n\nHer latest work, Shining Behind Shadows, marks a departure from her usual subject matter. — [2] —. Rather, the book highlights the lives and careers of twelve of her peers from Africa and Asia. The idea came to her three years ago at a conference in Chile. As Dr. Nujoma recalls, \"Listening to one speaker after another, I realized that many of my colleagues are from regions of the world, such as Africa and Asia, that were underrepresented.\" — [3] —.\n\nThe book is quite compelling. Dr. Nujoma brings to life the stories of how her peers were drawn to the profession and the effort they pour into their work. The book does, however, fall short in one respect: it provides insights only into the lives and careers of those working in the field of agricultural biotechnology. I would have welcomed the stories of those specializing in animal, marine, or medical biotechnology, too. — [4] —. Even so, Dr. Nujoma has again succeeded in creating a work that speaks to experts and laypeople alike.")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part7_g9 = [
        (169, "What is NOT indicated about Dr. Nujoma?",
         [("A", "She has visited Chile."), ("B", "She is a talented writer."), ("C", "She is based out of Namibia."), ("D", "She was interviewed by the London Daily Register.")], "D",
         "Nhà phê bình Paige Kinnock của tờ London Daily Register chỉ nhận định về sách của bà, không phỏng vấn trực tiếp bà."),
        (170, "What does Mr. Ellwood say about Shining Behind Shadows?",
         [("A", "It focuses on scientists from Africa and Asia."), ("B", "It was released three years ago."), ("C", "It details why Dr. Nujoma chose her career."), ("D", "It describes various branches of biotechnology.")], "A",
         "Cuốn sách tập trung vào cuộc đời và sự nghiệp của 12 nhà khoa học đến từ châu Phi và châu Á ('twelve of her peers from Africa and Asia')."),
        (171, "In which of the positions marked [1], [2], [3], and [4] does the following sentence best belong?\n\"Its focus is not on applying biotechnology in real-life situations.\"",
         [("A", "[1]"), ("B", "[2]"), ("C", "[3]"), ("D", "[4]")], "B",
         "Tại vị trí [2], câu trước nêu cuốn sách là bước rẽ khỏi chủ đề thường lệ ('departure from her usual subject matter') -> câu cần điền giải thích: trọng tâm không phải là ứng dụng công nghệ sinh học vào thực tiễn -> câu tiếp: 'Rather, the book highlights...'")
    ]
    for q_num, text, options, key, hint in part7_g9:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Group 172-175: Notice
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 172-175 refer to the following notice.\n\nKINGSTON-GARNET ISLAND PASSENGER FERRY SERVICE\nGeneral Information\n\nBeginning on 15 May, the Kingston-Garnet Island Passenger Ferry Service will resume service for eight weeks during the summer season. Ferries run daily every half hour from 7 a.m. to 8 p.m. The last ferry to Garnet Island will depart at 7:30 p.m. The last ferry from Garnet Island will leave at 8 p.m.\n\nBicycles are permitted on the passenger ferries. Bicyclists should arrive 30 minutes prior to departure and wait in the special bicycle lane to be loaded first. There are five racks that hold 50 bicycles on every passenger ferry.\n\nNo motorized vehicles are permitted on Garnet Island. Overnight parking is allowed in the main ferry terminal lot in Kingston. Rates are $5 per hour for up to 4 hours and a flat fee of $25 for four to 24 hours.\n\nVisit our Web site at www.kgferryservice.com for photographs of the ferry boats, a map of Garnet Island, lists of local attractions on the island, and information about peak-hour fare increases and group discounts.")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part7_g10 = [
        (172, "What is indicated about the ferry service?",
         [("A", "It is available only seasonally."), ("B", "Its boats were recently upgraded."), ("C", "It takes an hour to reach the island."), ("D", "It runs more frequently on weekends.")], "A",
         "Phà chỉ hoạt động 8 tuần trong mùa hè ('for eight weeks during the summer season')."),
        (173, "What is true about bicyclists on the ferries?",
         [("A", "They are last to board the boat."), ("B", "They must purchase a special ticket."), ("C", "They cannot travel on the 7:30 P.M. trip."), ("D", "They should arrive at the terminal early.")], "D",
         "Khách đi xe đạp được yêu cầu đến sớm 30 phút: 'Bicyclists should arrive 30 minutes prior to departure'."),
        (174, "What is not allowed on Garnet Island?",
         [("A", "Renting bicycles"), ("B", "Driving cars"), ("C", "Camping overnight"), ("D", "Taking photographs")], "B",
         "Quy định cấm phương tiện cơ giới: 'No motorized vehicles are permitted on Garnet Island.' -> không được lái ô tô."),
        (175, "What is indicated about the ferry tickets?",
         [("A", "They can be purchased at stores in Kingston."), ("B", "They are less expensive for children."), ("C", "They vary in price depending on the time of travel."), ("D", "They are more expensive if purchased on the boat.")], "C",
         "Có chính sách tăng giá vé vào giờ cao điểm: 'peak-hour fare increases' -> giá vé thay đổi tùy thời gian đi.")
    ]
    for q_num, text, options, key, hint in part7_g10:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Group 176-180: Receipt & Email
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 176-180 refer to the following receipt and e-mail.\n\nThank you for shopping at Green Stripe Press.\nOrder Number: GSP20896\nCustomer Information: Shoebox Mountain\nJason Ho <jasonho@shoeboxmountain.com>\nOrder Date: December 14 (PREPAID: online order)\nExpected Delivery: December 18–20\n\nQuantity | Item # | Item | Price\n1 | CAL201 | Complimentary Wild Animals Calendar | $ 0.00\n7 | ARB132 | Accounting Record Book ($19.99 each) | $139.93\nSubtotal: $139.93\nDiscount: $ 0.00\nTax @ 6%: $ 8.40\nShipping & Handling: $ 0.00\nTotal: $148.33\n\nThere is no charge for shipping and handling for corporate accounts. For questions regarding this order, please contact customerservice@greenstripepress.com.\n\n---\nTo: customerservice@greenstripepress.com\nFrom: Jason Ho <jasonho@shoeboxmountain.com>\nDate: December 18\nSubject: Order #GSP20896\n\nDear Green Stripe Press,\nI am writing regarding my most recent order (#GSP20896), which was delivered today. If you check my original order, you will see that I ordered six accounting record books. You sent us seven copies. I'd like to return the one I didn't order and have our corporate credit card refunded, together with the shipping cost the return will incur. Let me know how you would like me to proceed.\n\nOn a different note, congratulations on the calendar you included in my order! The photos are even more stunning than those in the Ancient Castles calendar you sent us last year. Some staff members saw mine and want copies of their own. Would you mind sending two more our way?\n\nBest,\nJason Ho, Owner\nShoebox Mountain")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part7_g11 = [
        (176, "Why was Mr. Ho not charged a shipping fee?",
         [("A", "He took advantage of a promotion."), ("B", "He made the purchase for his company."), ("C", "He picked up his order in person."), ("D", "He overpaid for shipping on a previous order.")], "B",
         "Hóa đơn ghi: 'There is no charge for shipping and handling for corporate accounts' (tài khoản doanh nghiệp mua cho công ty)."),
        (177, "What is true about Mr. Ho's order?",
         [("A", "It was damaged in transit."), ("B", "It included fragile items."), ("C", "It was paid for by check."), ("D", "It was delivered on time.")], "D",
         "Thời gian dự kiến nhận là 18–20 tháng 12, và thư phản hồi của ông vào ngày 18 tháng 12 nói hàng vừa giao hôm nay ('which was delivered today')."),
        (178, "According to the e-mail, what was Mr. Ho sent by mistake?",
         [("A", "A castle book he did not order"), ("B", "A calendar for last year"), ("C", "An incorrect refund check"), ("D", "An extra accounting book")], "D",
         "Ông chỉ đặt 6 quyển nhưng được giao 7 quyển sổ kế toán: 'sent us seven copies... return the one I didn't order'."),
        (179, "What is one reason Mr. Ho wrote the e-mail?",
         [("A", "To offer praise for an item"), ("B", "To complain about a price"), ("C", "To order some photo albums"), ("D", "To recommend a graphic designer")], "A",
         "Ông khen ngợi cuốn lịch tặng kèm rất đẹp: 'congratulations on the calendar you included in my order! The photos are even more stunning...'"),
        (180, "In the e-mail, the phrase \"our way\" in paragraph 2, line 4, is closest in meaning to",
         [("A", "in our style"), ("B", "to our address"), ("C", "at our expense"), ("D", "for our benefit")], "B",
         "Cụm từ 'sending two more our way' có nghĩa là gửi thêm hai cuốn tới địa chỉ của chúng tôi ('to our address').")
    ]
    for q_num, text, options, key, hint in part7_g11:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Group 181-185: Contact Form & Email
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 181-185 refer to the following contact form and e-mail.\n\nhttps://www.hilgravehistoricalcommission.org/contact_us\nContact the Hilgrave Historical Commission\nName: Roger Witmondt\nE-mail: roger@witmondtlocations.com\nComment:\nI work as a movie location scout. At this time, I need to find a location for a client's short film. The film will be set in the early 1930s, and I understand that Hilgrave has some interesting architecture from that era.\nCould you suggest places I should visit when I am in Hilgrave next month? Specifically, I am looking for a vacant building with an old-fashioned storefront that has distinctive details like stripe-patterned awnings and framed display windows. While having easy access to electricity would be ideal, it is not critical; my client has a generator that can be used if necessary. My client's aesthetic requirements are the most important consideration. Thank you for your assistance.\n\n---\nTo: roger@witmondtlocations.com\nFrom: brandi_schaertl@hilgravehistoricalcommission.org\nDate: 28 September\nSubject: The information you requested\nAttachment: Hilgrave Sites\n\nDear Mr. Witmondt,\nThank you for reaching out to the Hilgrave Historical Commission. The attached brochure lists the main areas of historical interest in Hilgrave. Some buildings may not meet all your requirements, but they might be worth considering.\nOne building that is not on the attached list is the old record store at 188 Main Street. It has been closed to business for the past ten years or so, but it seems to meet your criteria. The owner, Luke Nylund, is currently using it for storage; I would be happy to put you in touch with him. I doubt you will have trouble getting permission from him to use the space. While Hilgrave was once a bustling town, these days there are fewer businesses in operation. Many of the town's residents would appreciate the attention a short film might bring to the area.\n\nBest regards,\nBrandi Schaertl")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part7_g12 = [
        (181, "Why is Mr. Witmondt looking for a filming location in Hilgrave?",
         [("A", "It is known for its scenic mountain views."), ("B", "It is a short distance from his office."), ("C", "Its business district has appeared in other films."), ("D", "Its buildings represent a particular time period.")], "D",
         "Bộ phim lấy bối cảnh đầu thập niên 1930 và Hilgrave sở hữu kiến trúc từ thời kỳ lịch sử đó ('architecture from that era')."),
        (182, "On the contact form, the word \"critical\" in paragraph 2, line 4, is closest in meaning to",
         [("A", "judgmental"), ("B", "essential"), ("C", "sustainable"), ("D", "available")], "B",
         "'not critical' mang nghĩa không phải là điều kiện cốt yếu, bắt buộc (essential)."),
        (183, "What is implied about 188 Main Street?",
         [("A", "It was once used as a residence."), ("B", "It is frequently visited by tourists."), ("C", "It has several floors."), ("D", "It has decorative design features.")], "D",
         "Witmondt tìm kiếm cửa hàng có mái che kẻ sọc và cửa kính viền khung kiểu cổ điển; Brandi gợi ý số 188 Main Street vì đáp ứng các tiêu chuẩn thẩm mỹ này."),
        (184, "What does Ms. Schaertl offer to do?",
         [("A", "Advertise a forthcoming film"), ("B", "Arrange to have a building cleaned"), ("C", "Connect Mr. Witmondt with a building's owner"), ("D", "Help Mr. Witmondt acquire necessary permits from the town")], "C",
         "Bà đề nghị giới thiệu liên hệ với chủ tòa nhà Luke Nylund: 'I would be happy to put you in touch with him.'"),
        (185, "What does Ms. Schaertl suggest that Hilgrave needs?",
         [("A", "More publicity"), ("B", "Additional parking"), ("C", "A storage facility"), ("D", "A business directory")], "A",
         "Bà viết: 'Many of the town's residents would appreciate the attention a short film might bring to the area' (sự chú ý/quảng bá rộng rãi = publicity).")
    ]
    for q_num, text, options, key, hint in part7_g12:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Group 186-190: 2 Emails & Web page
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 186-190 refer to the following e-mails and Web page.\n\nTo: Cornelia Payne <cpayne@roughwing.co.uk>\nFrom: Pragya Mehta <pmehta@airsky.in>\nSubject: Lecture\nDate: 18 January\nAttachment: Notes\n\nDear Cornelia,\nI am sorry for the late notice, but I will not be able to join you for the lecture. My new position at the university in Mumbai requires me to remain on campus. Although I know you are fully prepared to deliver the lecture on your own, I have attached a copy of the notes I had prepared for the presentation. After you review them, let me know if there is anything else I can add. I had very much looked forward to traveling to Freeport, The Bahamas, for the first time and to seeing you again. I certainly miss working with you at the City University. Good luck with the lecture.\nPragya\n\n---\nTo: Ezra Halton <ehalton@ansonhouse.org>\nFrom: Cornelia Payne <cpayne@roughwing.co.uk>\nSubject: Information\nDate: 20 January\n\nDear Mr. Halton,\nI am pleased to confirm that 19 March works well for the lecture at the Anson House. As I had mentioned to you over the phone, I will be giving the lecture on my own. My publisher, Alphagamma Press, will be sending 30 copies of my most recent work, coauthored with Dr. Pragya Mehta, to your institution. You and I will both receive an e-mail confirmation from Alphagamma when the books ship. They should arrive at least a week before the lecture. I look forward to meeting you soon.\nCornelia Payne\n\n---\nhttps://www.theansonhouse.bs\nThe Anson House - Freeport, The Bahamas\nSpeaker Series\nTickets for lectures are $35 each or may be purchased at a reduced rate of $90 for the series of three. Complimentary refreshments will be served. Reservations are not required but are recommended.\n• 7 March, 7:00–8:30 P.M., Ms. Janelle Pears (historian and lifelong resident of Freeport).\n• 12 March, 7:00–8:30 P.M., Mr. Gregory Li (restoring historic gardens).\n• 19 March, 7:00–8:30 P.M., Dr. Cornelia Payne (discusses her book, The Transatlantic World of the Nineteenth Century, coauthored with Dr. Pragya Mehta. Dr. Payne has taught at the City University of Stoke-on-Trent in England for nearly 25 years).")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part7_g13 = [
        (186, "Why did Dr. Mehta send the e-mail to Dr. Payne?",
         [("A", "To cancel a vacation plan"), ("B", "To request lecture notes"), ("C", "To confirm a meeting"), ("D", "To offer an apology")], "D",
         "Tiến sĩ Mehta gửi email xin lỗi vì bận công việc mới ở Mumbai nên không thể đến Bahamas cùng thuyết trình ('I am sorry for the late notice...')."),
        (187, "According to the second e-mail, what did Dr. Payne do?",
         [("A", "Arrange a delivery"), ("B", "Give Mr. Halton a gift"), ("C", "Revise her presentation"), ("D", "Announce an address change")], "A",
         "Bà sắp xếp nhà xuất bản gửi 30 cuốn sách tới Anson House: 'My publisher, Alphagamma Press, will be sending 30 copies...'"),
        (188, "What is indicated about Alphagamma Press?",
         [("A", "It has headquarters in The Bahamas."), ("B", "It regularly ships materials to Anson House."), ("C", "It published The Transatlantic World of the Nineteenth Century."), ("D", "It is paying Dr. Payne's travel and accommodation expenses.")], "C",
         "Nhà xuất bản đã in cuốn sách mới nhất của Dr. Payne và Dr. Mehta: 'The Transatlantic World of the Nineteenth Century'."),
        (189, "Where were Dr. Mehta and Dr. Payne most likely colleagues?",
         [("A", "In Mumbai"), ("B", "In Freeport"), ("C", "In Stoke-on-Trent"), ("D", "In Charleston")], "C",
         "Dr. Mehta viết: 'I certainly miss working with you at the City University', và trên trang web ghi: 'City University of Stoke-on-Trent'."),
        (190, "What is indicated about the Speaker Series in the Web page?",
         [("A", "Discounts are unavailable."), ("B", "Reservations are optional."), ("C", "Events are held in the morning."), ("D", "Refreshments are not included.")], "B",
         "Trang web ghi rõ: 'Reservations are not required but are recommended' -> việc đặt chỗ trước là tùy chọn (optional).")
    ]
    for q_num, text, options, key, hint in part7_g13:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Group 191-195: Ad, Email, Press Release
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 191-195 refer to the following advertisement, e-mail, and press release.\n\nCRYN Group: We find the best employees for your company.\nPosted: November 25\nPosition title and codes:\n• Director of Operations, TL0015\n• Marketing Director, TL0023\n• Quality Control Director, TL0027\n• Director of Category Management, TL0045\nTo apply: Send your résumé to mdoro@cryngroup.ca.com with posting number 2098 in the subject line.\nCompany: Our client plans to chart an aggressive growth path in Latin America, where it plans to start operations next year. It is a well-established company widely known for marketing consumer-health products in North America, and more recently in Europe and Asia.\nQualifications: Candidates must have a formal business degree and a proven management record in at least one international setting. Experience with online sales and marketing preferred.\n\n---\nTo: Sven Arvidson <sarvidson@barkent.de.com>\nFrom: Maria Doro <mdoro@cryngroup.ca.com>\nRe: Posting number 2098\nDate: December 10\n\nDear Mr. Arvidson,\nThank you for submitting your résumé. I'd like to schedule a preliminary telephone conversation with you as soon as possible to determine whether you would be a good choice for our client. I would particularly like to discuss whether you would be prepared to work at our client's new facility overseas. Let me know if you are available for a 30-minute phone call sometime between 10:00 A.M. and 2:00 P.M. EST on Monday or Tuesday of next week. Please respond to me by e-mail at your earliest convenience.\nMaria Doro, CRYN Group\n\n---\nFor immediate release - Tayerson Ltd. Names New Director\nToronto, April 5—Tayerson Ltd. continues to staff its new international outpost, which opens one month from today. In this case, the lucky candidate is Sven Arvidson. \"Category Management is a relatively new area—and one that I'm sure I'll enjoy exploring and developing as director,\" Mr. Arvidson said. \"As a major online marketer of nutritional supplements, Tayerson is poised to lead the way toward better living while becoming a more vibrant, more profitable corporation.\"\nMr. Arvidson has held several key management positions throughout his career, most recently at Barkent Pharmaceuticals in Germany. He was also a partner at MSZ Consulting Group, where he provided marketing guidance to leading consumer-product companies in Canada and China.")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part7_g14 = [
        (191, "According to the advertisement, what qualification is not required of applicants?",
         [("A", "A degree in business"), ("B", "Previous employment in a managerial position"), ("C", "Experience working overseas"), ("D", "Online sales and marketing experience")], "D",
         "Yêu cầu 'Experience with online sales and marketing preferred' chỉ là ưu tiên thêm, không bắt buộc (not required)."),
        (192, "What most likely is the location of the position for which Mr. Arvidson applied?",
         [("A", "Latin America"), ("B", "North America"), ("C", "Europe"), ("D", "Asia")], "A",
         "Khách hàng tuyển dụng để chuẩn bị mở cơ sở mới tại khu vực Mỹ Latinh ('plans to start operations next year in Latin America')."),
        (193, "What does Ms. Doro ask Mr. Arvidson to provide in his reply?",
         [("A", "Information on his leadership skills"), ("B", "His availability for an interview"), ("C", "A current résumé"), ("D", "The names of two references")], "B",
         "Bà Doro đề nghị ông cung cấp thời gian rảnh để phỏng vấn qua điện thoại: 'Let me know if you are available for a 30-minute phone call...'"),
        (194, "What job code did Mr. Arvidson most likely reference in his application?",
         [("A", "TL0015"), ("B", "TL0023"), ("C", "TL0027"), ("D", "TL0045")], "D",
         "Ông được bổ nhiệm làm Giám đốc Quản lý Ngành hàng ('Director of Category Management'), có mã tuyển dụng là TL0045."),
        (195, "According to the press release, what does Tayerson Ltd. sell?",
         [("A", "Medical devices"), ("B", "Accounting software"), ("C", "Nutritional supplements"), ("D", "Exercise equipment")], "C",
         "Thông cáo báo chí nêu rõ: 'online marketer of nutritional supplements' -> bán thực phẩm chức năng bổ sung dinh dưỡng.")
    ]
    for q_num, text, options, key, hint in part7_g14:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    # Group 196-200: Brochure, Email, Schedule
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 196-200 refer to the following brochure, e-mail, and schedule.\n\nFGJ's Business Expert Series: Prospect Data Purchasing\n45-Minute Webinar | 11 May, 2:30 P.M.\nEvery department in an organisation needs information to make business decisions. Marketing professionals in particular rely on accurate data about potential customers to be successful. In this Webinar, Briana Carrera, chief financial officer of Pile One Market Data, explains how to avoid the pitfalls of acquiring incomplete data and provides tips on what to ask your prospective data provider to ensure that your next batch of data leads to the results you want.\n\n---\n*E-mail*\nTo: Sandra Lescure\nFrom: Gino Stelleti\nDate: 12 May\nSubject: FGJ's Webinar\n\nDear Sandra,\nYou had asked me to give you a report on the Webinar I attended yesterday about acquiring data. To be honest, even though the facilitator was clearly knowledgeable, she did not tell me anything I didn't already know. I was hoping that she would show us the differences between the many types of databases out there, but she never touched on that topic. I am not so sure all these Webinars are useful; in the end, they always seem to be steering us to purchase from one particular company or another—in this case it was Pile One. Having said that, the new schedule just came out, and there is another Webinar coming up soon that I am hoping will provide the information I want about how to store and sort data. I have already signed up for it; I will keep you posted about what I learn.\nGino\n\n---\nUpcoming 45-Minute Webinars in FGJ's Business Expert Series\n• Basics of Market Research | 1 June, 9:30 A.M. | Ed Quinones | €12.00\n• Tips for Accelerating Sales | 12 June, 9:30 A.M. | Cameron Stone | €16.00\n• What is Market Automation? | 5 July, 11:00 A.M. | Ed Quinones | €12.00\n• Choosing the Right Database | 17 July, 2:00 P.M. | Selina Tucci | €12.00\n\nThe Business Expert Series is organised by Tuyet Nguyen. Please address any questions to her at tnguyen@fgj.org. To access Webinar content, attendance is required as Webinars are not recorded for later viewing.")
    doc.add_paragraph("[PASSAGE_END]")
    doc.add_paragraph("[GROUP_SHUFFLE_Q: FALSE]")
    doc.add_paragraph()

    part7_g15 = [
        (196, "According to the brochure, who would benefit most from the Prospect Data Purchasing Webinar?",
         [("A", "A financial analyst"), ("B", "A marketing manager"), ("C", "A customer-service representative"), ("D", "An information-technology specialist")], "B",
         "Tờ rơi ghi: 'Marketing professionals in particular rely on accurate data...' -> chuyên viên marketing hưởng lợi nhiều nhất."),
        (197, "What was Mr. Stelleti's complaint about the Webinar?",
         [("A", "He did not like the format."), ("B", "He did not learn anything new."), ("C", "He had difficulty hearing everything."), ("D", "He found the topic to be too complex.")], "B",
         "Ông phàn nàn: 'she did not tell me anything I didn't already know' -> không học thêm được điều gì mới."),
        (198, "What does Mr. Stelleti suggest about Ms. Carrera?",
         [("A", "She recently joined a new company."), ("B", "She has experience managing databases."), ("C", "She had asked him to sign up for the Webinar."), ("D", "She tried to sell her company's services.")], "D",
         "Gino nhận xét: 'steering us to purchase from one particular company or another—in this case it was Pile One' (nơi cô làm Giám đốc tài chính)."),
        (199, "When will Mr. Stelleti most likely attend another Webinar?",
         [("A", "On June 1"), ("B", "On June 12"), ("C", "On July 5"), ("D", "On July 17")], "D",
         "Ông muốn biết về 'how to store and sort data', tương ứng với buổi 'Choosing the Right Database' diễn ra vào ngày 17 tháng 7."),
        (200, "According to the schedule, what do all of the Webinars have in common?",
         [("A", "They cost the same amount."), ("B", "They are held in the morning."), ("C", "They last the same amount of time."), ("D", "They are recorded for future playback.")], "C",
         "Tiêu đề lịch thông báo: 'Upcoming 45-Minute Webinars in FGJ's Business Expert Series' -> tất cả các buổi đều kéo dài 45 phút.")
    ]
    for q_num, text, options, key, hint in part7_g15:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for opt_label, opt_text in options:
            doc.add_paragraph(f"[{opt_label}] {opt_text}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph()

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Generated successfully: {output_path}")

if __name__ == "__main__":
    import sys
    out = sys.argv[1] if len(sys.argv) > 1 else "ETS_Test_09_Tagged.docx"
    create_tagged_exam_docx(out)
