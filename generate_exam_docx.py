import docx

def create_toeic_reading_test_09():
    doc = docx.Document()

    # Tiêu đề và Thời gian làm bài
    doc.add_paragraph("[EXAM_TITLE: ETS TOEIC Reading Test 09]")
    doc.add_paragraph("[DURATION: 75]")
    doc.add_paragraph("")

    # ==========================================
    # PART 5: CÂU HỎI ĐỘC LẬP (101 - 130)
    # ==========================================
    doc.add_paragraph("[PART: 5]")

    part5_data = [
        (101, "Every batch of sauce at Generita's Bistro is processed meticulously by ------- expert chefs.",
         ["they", "their", "them", "themselves"], "B", "Đứng trước cụm danh từ 'expert chefs' cần một tính từ sở hữu. Chọn their."),
        (102, "Computerization of medical records ------- increases a physician's ability to diagnose and treat patients.",
         ["great", "greatly", "greatness", "greatest"], "B", "Đứng trước động từ thường 'increases' cần một phó từ (trạng từ) bổ nghĩa. Chọn greatly."),
        (103, "At Rojelle's Fine Dining, we use the freshest ------- available to make our salads.",
         ["applications", "subjects", "ingredients", "factors"], "C", "Dựa vào ngữ cảnh món ăn/salad ('make our salads'), từ phù hợp nhất là 'ingredients' (nguyên liệu)."),
        (104, "Professor Benguigui will present his paper ------- the natural history conference tomorrow.",
         ["by", "at", "of", "on"], "B", "Giới từ 'at' dùng để chỉ địa điểm, sự kiện cụ thể: 'at the natural history conference'."),
        (105, "The chefs are still waiting for a ------- date for the new convection oven that was ordered two weeks ago.",
         ["shipper", "ships", "shipments", "shipping"], "D", "Cụm danh từ ghép 'shipping date' có nghĩa là ngày giao hàng."),
        (106, "Mr. Kang works ------- with our internal team members as well as various regional sales representatives.",
         ["mildly", "nearly", "closely", "narrowly"], "C", "Cụm 'work closely with' có nghĩa là hợp tác chặt chẽ với ai đó."),
        (107, "Fashion designer Hye-Ja Pak knows ------- to update her line in response to changing tastes.",
         ["and", "when", "need", "for"], "B", "Cấu trúc từ để hỏi kết hợp động từ nguyên mẫu: 'knows when to update' (biết khi nào nên cập nhật)."),
        (108, "The samba class was so well ------- that the Yulara Dance School decided to make the course a permanent offering.",
         ["attended", "educated", "gathered", "protected"], "A", "Cụm 'well attended' nghĩa là có đông người tham dự/tham gia."),
        (109, "The factory will be situated away ------- the city's residential area to reduce complaints about noise and emissions.",
         ["from", "about", "with", "out"], "A", "Cụm cố định 'away from': cách xa khỏi khu vực dân cư."),
        (110, "Ms. Kuramoto selected the most ------- mailing option available.",
         ["economical", "economy", "economize", "economized"], "A", "Sau 'the most' và trước danh từ 'mailing option' cần một tính từ dài: 'economical' (tiết kiệm)."),
        (111, "Companies without information technology specialists can ------- on Vyber Software Advisers for assistance with online services.",
         ["reliable", "reliably", "rely", "relying"], "C", "Sau modal verb 'can' cần một động từ nguyên thể không chia: 'rely on' (trông cậy vào)."),
        (112, "Next year, our team will have a new task, ------- is to review design portfolios.",
         ["although", "which", "after", "because"], "B", "Đại từ quan hệ 'which' thay thế cho mệnh đề hoặc danh từ phía trước ('a new task') làm chủ ngữ."),
        (113, "Cook the meat for 30 minutes to ensure ------- readiness to be eaten.",
         ["both", "this", "its", "that"], "C", "Cần một tính từ sở hữu bổ nghĩa cho danh từ 'readiness': 'its readiness' (độ chín của nó)."),
        (114, "Amoxitron's research team will hire ------- interns to assist with laboratory duties.",
         ["given", "several", "whole", "natural"], "B", "Từ 'several' (một vài) đứng trước danh từ số nhiều đếm được 'interns'."),
        (115, "------- the new acai juice blend has proved so popular, we should move quickly to increase our production volume.",
         ["If", "Whether", "Since", "Unless"], "C", "'Since' ở đây đóng vai trò liên từ chỉ nguyên nhân (Bởi vì/Vì... nên chúng ta phải hành động nhanh chóng)."),
        (116, "New emissions standards have forced Rider Auto to modify the process of engine -------.",
         ["construction", "constructed", "constructive", "construct"], "A", "Sau giới từ 'of' cần một danh từ: cụm 'engine construction' (việc chế tạo động cơ)."),
        (117, "Maki Kayano's book offers techniques for ------- business tasks with speed and precision.",
         ["executing", "equipping", "returning", "involving"], "A", "Sau giới từ 'for' là V-ing: 'executing business tasks' mang nghĩa thực thi/hoàn thành nhiệm vụ kinh doanh."),
        (118, "To control costs, updated credit card readers will be installed in branch stores -------.",
         ["gradual", "gradually", "more gradual", "most gradual"], "B", "Đứng ở cuối câu bổ nghĩa cho động từ bị động 'will be installed', dùng trạng từ 'gradually' (từng bước/dần dần)."),
        (119, "In this quarter, the Montel Beverage Company is expecting sales ------- £160,000 and £180,000.",
         ["without", "among", "throughout", "between"], "D", "Cấu trúc quen thuộc: 'between A and B' (ở khoảng giữa A và B)."),
        (120, "Because Mount Akoyola is so challenging for climbers, ------- have reached its peak.",
         ["any", "either", "other", "few"], "D", "'Few' đóng vai trò đại từ mang nghĩa 'rất ít người' có thể chạm tới đỉnh."),
        (121, "King Street Bridge will be closed in the month of September ------- repair work.",
         ["because of", "so that", "as if", "rather than"], "A", "Phía sau là cụm danh từ 'repair work', dùng 'because of' (bởi vì công việc sửa chữa)."),
        (122, "Ms. Taniguchi's supervisor commended her for negotiating ------- with Furuyama Corporation.",
         ["effective", "effecting", "effected", "effectively"], "D", "Bổ nghĩa cho động từ/V-ing 'negotiating' cần trạng từ 'effectively' (đàm phán một cách hiệu quả)."),
        (123, "Staff members may reserve the conference room ------- they need it.",
         ["somewhere", "whatever", "everything", "anytime"], "D", "'Anytime' đóng vai trò liên từ mang nghĩa bất kỳ lúc nào / bất cứ khi nào họ cần."),
        (124, "The public relations director must have a high level of ------- in English and Spanish.",
         ["proficiency", "advancement", "routine", "strength"], "A", "Cụm danh từ 'proficiency in' có nghĩa là sự thành thạo/thông thạo một ngôn ngữ."),
        (125, "Patterson Products ------- seeks innovative ways of meeting changing consumer demand.",
         ["dually", "favorably", "continually", "generically"], "C", "Trạng từ 'continually' (liên tục/không ngừng) bổ nghĩa cho động từ 'seeks'."),
        (126, "Mr. Volante is working at home tomorrow so he ------- the technical report without any distractions.",
         ["can finish", "would finish", "finished", "has been finishing"], "A", "Mệnh đề chỉ mục đích với 'so': 'can finish' (để anh ấy có thể hoàn thành báo cáo)."),
        (127, "All of the billing procedures remain the same, ------- new payment codes need to be inserted into the invoice documents.",
         ["in order that", "during", "across from", "except that"], "D", "Liên từ 'except that' có nghĩa là 'ngoại trừ việc/chỉ trừ việc... còn lại giữ nguyên'."),
        (128, "The lightweight design of the new sedan is ------- because it can cause the vehicle to slide on icy roads.",
         ["possible", "mechanical", "questionable", "multiple"], "C", "Tính từ 'questionable' nghĩa là đáng nghi ngại / còn nhiều nghi vấn do có thể gây trượt xe."),
        (129, "When the bank president retired, the common ------- was that the vice president would take over.",
         ["assume", "assumption", "assuming", "assumable"], "B", "Sau tính từ 'common' cần một danh từ đóng vai trò chủ ngữ: 'the common assumption' (giả định phổ biến)."),
        (130, "West Bengali Airlines ------- fees for oversized items that are still within weight limitations.",
         ["invites", "cooperates", "transports", "waives"], "D", "Động từ 'waives' (miễn trừ): hãng hàng không miễn trừ phí đối với các kiện hàng quá khổ nhưng vẫn nằm trong giới hạn trọng lượng.")
    ]

    for q_num, text, opts, key, hint in part5_data:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: TRUE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for idx, opt_label in enumerate(["A", "B", "C", "D"]):
            doc.add_paragraph(f"[{opt_label}] {opts[idx]}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph("")

    # ==========================================
    # PART 6: CÂU HỎI THEO ĐOẠN VĂN (131 - 146)
    # ==========================================
    doc.add_paragraph("[PART: 6]")

    # Nhóm 1: 131 - 134
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 131-134 refer to the following e-mail.\n"
                      "To: pmendoza@factmail.co\n"
                      "From: recruiting@analystsassoc.org\n"
                      "Date: May 2\n"
                      "Subject: Joining ASA\n\n"
                      "Dear Mr. Mendoza,\n"
                      "Thank you for expressing your interest in our organization during the recent ---[131]---. It was a pleasure "
                      "meeting you at the ASA booth during the Weber Information Systems Convention. As you may recall from our conversation, "
                      "we discussed how ASA membership ---[132]--- your career through networking opportunities as well as the industry insights "
                      "offered in our monthly newsletter. The normal fee for members is just $120 a year; however, we are currently offering a new-member "
                      "discount. ---[133]---.\n"
                      "If you are still interested in joining, please reply to ---[134]--- with your mailing address. I will then forward you an application packet.\n\n"
                      "Sincerely,\nAshlee Loren, President\nAssociation of Systems Analysts")
    doc.add_paragraph("[PASSAGE_END]")

    p6_group1 = [
        (131, "Chọn từ phù hợp vào chỗ trống [131]", ["election", "broadcast", "conference", "performance"], "C", "Dựa vào ngữ cảnh nhắc đến 'Weber Information Systems Convention' (hội nghị), từ đồng nghĩa là conference."),
        (132, "Chọn dạng động từ phù hợp vào chỗ trống [132]", ["can benefit", "is benefiting", "has benefited", "will have benefited"], "A", "'can benefit' diễn tả khả năng mang lại lợi ích nghề nghiệp cho thành viên."),
        (133, "Chọn câu phù hợp điền vào vị trí [133]", [
            "Discounts on rental cars are included in the annual membership.",
            "Our jobs board is quite comprehensive.",
            "We are the first organization of our kind.",
            "This month you can join for just half the usual rate."
        ], "D", "Câu trước nói về việc giảm giá cho hội viên mới ('offering a new-member discount'), nên câu phù hợp nhất là gia nhập với nửa giá thông thường."),
        (134, "Chọn đại từ phù hợp vào chỗ trống [134]", ["it", "me", "them", "anyone"], "B", "Người viết thư (Ashlee Loren) yêu cầu 'reply to me' (phản hồi lại cho tôi).")
    ]
    for q_num, text, opts, key, hint in p6_group1:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for idx, opt_label in enumerate(["A", "B", "C", "D"]):
            doc.add_paragraph(f"[{opt_label}] {opts[idx]}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph("")

    # Nhóm 2: 135 - 138
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 135-138 refer to the following e-mail.\n"
                      "To: Sam Heinz\n"
                      "From: Northways Professional Development\n"
                      "Sent: April 20\n"
                      "Subject: Workshop 4/28-4/30\n\n"
                      "Dear Workshop Participants,\n"
                      "We look forward to seeing you at the digital storytelling workshop. On the first day, when you enter the Albin College campus, "
                      "attendants will ---[135]--- you to lot 43 and the Toteman Building. We will begin each day in conference room 9. "
                      "Coffee, tea, snacks, and fruit ---[136]--- in the mornings. Lunch will be sandwiches and salads from Black Horse Restaurant. ---[137]---.\n"
                      "In the afternoons, we will be working in the computer lab. We suggest that you assemble some images that you would like to use for your project. "
                      "It will be ---[138]--- for you to have them saved on a digital storage device beforehand. If you have any questions, please e-mail us.\n\n"
                      "Very Best,\nGina Kapuski")
    doc.add_paragraph("[PASSAGE_END]")

    p6_group2 = [
        (135, "Chọn từ phù hợp vào chỗ trống [135]", ["offer", "direct", "pass", "instruct"], "B", "'direct you to lot 43' nghĩa là người hướng dẫn sẽ chỉ đường cho bạn đến bãi đỗ xe."),
        (136, "Chọn dạng động từ phù hợp vào chỗ trống [136]", ["will be provided", "were provided", "providers", "are providing"], "A", "Sự việc sẽ diễn ra trong tương lai ở dạng bị động: 'will be provided' (sẽ được chuẩn bị sẵn)."),
        (137, "Chọn câu phù hợp điền vào vị trí [137]", [
            "Some prior experience working with digital files is assumed.",
            "Please do not bring any copyrighted material to the workshop.",
            "Please let us know if you have any dietary restrictions.",
            "There is a one-time parking charge of fifteen dollars."
        ], "C", "Câu trước nói về bữa trưa ('Lunch will be sandwiches...'), nên câu tiếp theo liên quan đến chế độ ăn uống: 'dietary restrictions'."),
        (138, "Chọn tính từ phù hợp vào chỗ trống [138]", ["useful", "surprised", "difficult", "amazing"], "A", "'It will be useful for you...' nghĩa là sẽ rất hữu ích nếu bạn lưu trước ảnh vào thiết bị lưu trữ số.")
    ]
    for q_num, text, opts, key, hint in p6_group2:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for idx, opt_label in enumerate(["A", "B", "C", "D"]):
            doc.add_paragraph(f"[{opt_label}] {opts[idx]}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph("")

    # Nhóm 3: 139 - 142
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 139-142 refer to the following article.\n\n"
                      "Venley Foods Responsive to Changing Consumers\n"
                      "Who cares where the tomatoes in your salad actually came from? ---[139]---, an increasing number of people do, "
                      "according to a study conducted by the Consumer Group. In fact, the study shows that many consumers would pay an average of 10 percent "
                      "more when given the exact source of a fresh food product. ---[140]---. Some grocery stores, such as Venley Foods in Boston, "
                      "have taken advantage of the trend and used it to ---[141]--- implement branding and marketing. \"If we can tell a story about our product,\" "
                      "says Venley Foods CEO Minji Kim, \"then we've added ---[142]--- in the minds of consumers.\"")
    doc.add_paragraph("[PASSAGE_END]")

    p6_group3 = [
        (139, "Chọn phó từ phù hợp vào chỗ trống [139]", ["Carefully", "Apparently", "Formerly", "Rarely"], "B", "'Apparently' (rõ ràng là/hóa ra là) dùng để mở đầu câu phản hồi cho câu hỏi trước đó."),
        (140, "Chọn câu phù hợp điền vào vị trí [140]", [
            "Fresh food can be refrigerated for up to two days.",
            "Many grocery stores have been extending their hours.",
            "Most studies are published in consumer magazines.",
            "The number increases to 20 percent in large cities."
        ], "D", "Câu trước nói về việc trả thêm 10%, câu sau tiếp nối số liệu thống kê: tỷ lệ này tăng lên 20% tại các thành phố lớn."),
        (141, "Chọn từ loại phù hợp vào chỗ trống [141]", ["smarts", "smartly", "smarter", "smartness"], "B", "Bổ nghĩa cho động từ 'implement' cần một trạng từ: 'smartly implement'."),
        (142, "Chọn danh từ phù hợp vào chỗ trống [142]", ["value", "time", "obstacles", "bonus"], "A", "'add value' là cụm từ quen thuộc chỉ việc gia tăng giá trị cho sản phẩm.")
    ]
    for q_num, text, opts, key, hint in p6_group3:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for idx, opt_label in enumerate(["A", "B", "C", "D"]):
            doc.add_paragraph(f"[{opt_label}] {opts[idx]}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph("")

    # Nhóm 4: 143 - 146
    doc.add_paragraph("[PASSAGE_START]")
    doc.add_paragraph("Questions 143-146 refer to the following article.\n\n"
                      "Nylobe, Inc., Announces Newest Development Project\n"
                      "SEATTLE (May 1) - Researchers at technology firm Nylobe, Inc., are working to develop a sensor capable of detecting "
                      "corrosion caused by environmental exposure. Corrosion is a major contributor to ---[143]--- losses in the aircraft industry each year. "
                      "\"This will be a major ---[144]--- for commercial airline fleets,\" says Mel Laveau, Nylobe's CEO. \"The sensor will decrease both labor "
                      "and maintenance costs without being too expensive.\"\n"
                      "According to Ms. Laveau, the sensor will work by detecting corrosion in its early stages, when the problem can be corrected simply "
                      "by removing the corroded material. ---[145]---. In the structure of large aircraft, some critical joints can be particularly "
                      "susceptible to corrosion. ---[146]---, the sensor can be used to inspect these areas and then target the most likely areas of concern.")
    doc.add_paragraph("[PASSAGE_END]")

    p6_group4 = [
        (143, "Chọn từ loại phù hợp vào chỗ trống [143]", ["financially", "financed", "financial", "finances"], "C", "Đứng trước danh từ 'losses' cần một tính từ: 'financial losses' (tổn thất tài chính)."),
        (144, "Chọn danh từ phù hợp vào chỗ trống [144]", ["balance", "examination", "expectation", "asset"], "D", "'a major asset' mang ý nghĩa là một tài sản/công cụ đắc lực, vô cùng hữu ích."),
        (145, "Chọn câu phù hợp điền vào vị trí [145]", [
            "This will reduce the need for making expensive structural repairs.",
            "The parts have all been replaced with higher quality materials.",
            "The next stage of the project involves scanning the affected areas.",
            "Its style and sleek design made it popular with the public."
        ], "A", "Việc phát hiện ăn mòn sớm sẽ giúp giảm thiểu nhu cầu sửa chữa kết cấu tốn kém."),
        (146, "Chọn liên từ phù hợp vào chỗ trống [146]", ["Meanwhile", "Similarly", "Otherwise", "Fortunately"], "D", "'Fortunately' (May mắn thay), cảm biến có thể được dùng để kiểm tra các vị trí hiểm hóc này.")
    ]
    for q_num, text, opts, key, hint in p6_group4:
        doc.add_paragraph(f"[Q: {q_num}] {text}")
        doc.add_paragraph("[SHUFFLE_Q: FALSE]")
        doc.add_paragraph("[SHUFFLE_A: TRUE]")
        for idx, opt_label in enumerate(["A", "B", "C", "D"]):
            doc.add_paragraph(f"[{opt_label}] {opts[idx]}")
        doc.add_paragraph(f"[KEY: {key}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[END_Q]")
        doc.add_paragraph("")

    # ==========================================
    # PART 7: ĐOẠN ĐƠN VÀ ĐOẠN KÉP/BA (147 - 200)
    # ==========================================
    doc.add_paragraph("[PART: 7]")

    part7_sections = [
        # 147-148
        ("Questions 147-148 refer to the following invoice.\n\n"
         "Invoice 3987\nShawqi Office Services, Dubai, UAE\n18 October\n"
         "Jenkins Press, P.O. Box 2291, Dubai, UAE\n\n"
         "Service: On 14 October, replaced bulb and repaired paper tray on copier per call received on 12 October. "
         "Replaced copy ink in two machines. Performed routine yearly maintenance on five copiers per existing service contract.\n"
         "Labor cost: AED 330.00 | Paper tray: AED 50.00 | Bulb: AED 30.00 | Copy ink: AED 220.00\n"
         "Total: AED 630.00\nTotal amount must be received by 31 October. Thank you for your business!",
         [
             (147, "What is indicated about Jenkins Press?",
              ["It has several offices around the world.", "Its copiers get checked every year.", "Its office equipment is outdated.", "It is a new customer of Shawqi Office Services."],
              "B", "Hóa đơn ghi: 'Performed routine yearly maintenance on five copiers' -> máy photocopy được kiểm tra/bảo dưỡng hàng năm."),
             (148, "When is payment due?",
              ["October 12", "October 14", "October 18", "October 31"],
              "D", "Dòng cuối ghi: 'Total amount must be received by 31 October.'")
         ]),

        # 149-151
        ("Questions 149-151 refer to the following advertisement.\n\n"
         "Harbour View Apartment - Porthmadog, Wales\n"
         "This one-bedroom apartment is perfect for a holiday escape! Located in a quiet area, it boasts a patio with a lovely view of the harbour. "
         "Recently renovated, the unit includes an eat-in kitchen with stove, refrigerator, microwave, and coffeepot; bathroom with walk-in shower; "
         "and a living room with a large-screen TV. Other amenities of the property include:\n"
         "• Short distance to restaurants and shops\n• Five-minute walk to the beach\n"
         "• Public gardens and historic sites within a 20-minute drive\n• Heat and electricity included\n"
         "• Towels and bed linens provided on-site\n• Daily cleaning service available (extra fee)\n"
         "• Wireless Internet access (extra fee)\n\n"
         "Reserve this lovely gem now! Signing a contract by March 30 will reduce the rental cost by 10 percent. "
         "To sign a contract, contact Dylan Barrett at dbarrett@telarentals.co.uk.",
         [
             (149, "Who would the advertisement most likely interest?",
              ["Business travelers", "Residents of Porthmadog", "Property investors", "Short-term vacationers"],
              "D", "Đoạn văn mở đầu nêu: 'perfect for a holiday escape' -> đối tượng là khách du lịch ngắn ngày."),
             (150, "What is indicated about the rental fee?",
              ["It includes tours of historic places.", "It does not cover all of the apartment's features.", "It includes vouchers to use at local restaurants.", "It requires a minimum 30 percent deposit in advance."],
              "B", "Các tiện ích như dịch vụ dọn phòng hàng ngày và wifi phải trả thêm phí ('extra fee')."),
             (151, "Why should an individual contact Mr. Barrett by March 30?",
              ["To get a discount on rent", "To schedule apartment renovations", "To rent the last available property", "To sell the property before the end of the season"],
              "A", "Ký hợp đồng trước ngày 30/3 sẽ được giảm giá 10% ('reduce the rental cost by 10 percent').")
         ]),

        # 152-153
        ("Questions 152-153 refer to the following text-message chain.\n\n"
         "Jason Salter (1:45 P.M.): Liz, the meeting is starting in fifteen minutes. Where are you?\n"
         "Liz Ortiz (1:47 P.M.): The train has been stopped on the tracks for a while. There seems to be some kind of problem. I still hope to make the meeting.\n"
         "Jason Salter (1:50 P.M.): Okay. I'll save you a seat.\n"
         "Liz Ortiz (1:59 P.M.): There's just been an announcement. There's a disabled train up ahead. It'll be awhile.\n"
         "Jason Salter (2:00 P.M.): Don't worry. If there are any questions for our department, I'll handle them.\n"
         "Liz Ortiz (2:01 P.M.): Thanks. I'll call you later.",
         [
             (152, "What is suggested about Mr. Salter?",
              ["He is leading the meeting.", "He commutes to work by train.", "He has many questions for Ms. Ortiz.", "He works in the same department as Ms. Ortiz."],
              "D", "Jason nói: 'If there are any questions for our department, I'll handle them' -> hai người làm chung một phòng ban."),
             (153, "At 1:59 P.M., what does Ms. Ortiz imply when she writes, 'It'll be awhile'?",
              ["The meeting is running late.", "She is still preparing her notes.", "She will likely miss the meeting.", "She has not boarded the train yet."],
              "C", "Tàu bị kẹt và mất nhiều thời gian nên cô ấy có thể sẽ bỏ lỡ cuộc họp.")
         ]),

        # 154-155
        ("Questions 154-155 refer to the following e-mail.\n\n"
         "From: tbogosian@sardhasconvention.com\nTo: pradalaily@dmcv.com\nDate: 22 March\nSubject: Your reservation - Prada party\n\n"
         "Dear Ms. Prada,\n"
         "Thank you for choosing Sardha's Banquet Hall to host your event on 30 August. Unfortunately, West Hall will be under renovation in August, "
         "and we will be unable to accommodate your group in that room. However, your party will fit comfortably in East Hall. "
         "This room features floor-to-ceiling windows and a full view of the river. The 15 percent deposit we have received will hold the reservation.\n"
         "As indicated when you completed your request through our Web site, your total price of $1,600 includes a full-course meal that includes "
         "appetizers, soup or salad, main course, and dessert. Attached to this e-mail please find a copy of the menu presented to guests with options for appetizer and main course.\n\n"
         "Tomas Bogosian, General Manager",
         [
             (154, "What is one purpose of the e-mail?",
              ["To cancel a reservation", "To apologize for an error", "To indicate a room change", "To request an extra deposit"],
              "C", "West Hall đang sửa chữa nên chuyển sang East Hall -> thông báo đổi phòng."),
             (155, "What is suggested about Ms. Prada?",
              ["She has paid her bill in full.", "She eats at Sardha's regularly.", "She works in the food industry.", "She made the reservation online."],
              "D", "Email nêu: 'when you completed your request through our Web site' -> cô ấy đặt trước trên trang web.")
         ]),

        # 156-157
        ("Questions 156-157 refer to the following memo.\n\n"
         "MEMO\nTo: All Employees\nFrom: IT Supervisor\nSubject: OS Update\nDate: 27 May\n\n"
         "All company computers in the Melbourne and Victoria offices require an update to the operating system. "
         "This update will prepare our computers for the new version of our accounting software, which should arrive in the first week of July. "
         "Technicians will install the new operating system beginning on 3 June. We expect the process to be completed around 16 June, "
         "providing time to work out any bugs with the operating system before the accounting software is uploaded.\n"
         "Completing the entire process will require you to log in so that we can verify the system is working as anticipated. "
         "Therefore, if you plan to go on holiday during this period, please inform IT Support immediately of the dates you will be out "
         "so that we can plan an alternate date to accommodate your schedule.\n"
         "If you have any questions, please contact IT Support staff at extension 48.",
         [
             (156, "When is the installation of the operating system expected to be finished?",
              ["In the first week of June", "In the middle of June", "At the end of June", "In the first week of July"],
              "B", "Văn bản nêu: 'completed around 16 June' -> vào khoảng giữa tháng 6."),
             (157, "Who is asked to contact IT Support?",
              ["Employees who do not need the accounting software", "Accountants already using the new operating system", "Staff who work outside the offices", "Employees who are taking time off"],
              "D", "Ai dự định đi nghỉ/nghỉ phép ('plan to go on holiday') thì cần báo cho IT Support.")
         ]),

        # 158-160
        ("Questions 158-160 refer to the following letter.\n\n"
         "Clarke-Ellis Construction\n#20 Murphy Industrial Park, St. Michael BB23028\n4 June\n"
         "Ida Gutierrez, Darling Cove Inn, Mango Drive, Folkestone BB24017\n\n"
         "Dear Ms. Gutierrez,\n"
         "Thank you for contacting Clarke-Ellis Construction for your roofing project. - [1] -. "
         "After inspecting the property, I have confirmed that the inn's main roof is in good condition and requires no repairs at this time. "
         "However, the roof of the inn's porch appears to be at least twenty years old and is worn beyond repair. - [2] -. "
         "Clarke-Ellis Construction can remove and dispose of the existing porch roof and install a new one. The replacement will be comparable "
         "to the quality, style, and colour of the inn's main roof. We will use only commercial-grade leak barriers, insulation, and shingles "
         "manufactured by West Indies Weatherproofing, Inc. The estimated total cost, inclusive of labour and materials, is $3,260. - [3] -. "
         "Additional fees would apply should you want us to make other improvements, such as painting the porch or replacing porch screens. - [4] -.\n"
         "Please call me at the number above to discuss scheduling or any questions you have about the work. I hope to hear from you soon.\n\n"
         "Sincerely,\nGrayson Clarke, Co-owner, Clarke-Ellis Construction",
         [
             (158, "Why did Mr. Clarke write the letter?",
              ["To provide a work proposal", "To ask for a project extension", "To request an inspection report", "To submit a revised cost estimate"],
              "A", "Thư cung cấp đề xuất phương án và chi phí sửa mái hiên ('work proposal')."),
             (159, "What is indicated about West Indies Weatherproofing, Inc.?",
              ["It is installing a new porch at Darling Cove Inn.", "It provides painting services.", "It produces roofing materials.", "It is owned by Clarke-Ellis Construction."],
              "C", "Thư nêu: 'shingles manufactured by West Indies Weatherproofing, Inc.' -> công ty này sản xuất vật liệu lợp mái."),
             (160, "In which of the positions marked [1], [2], [3], and [4] does the following sentence best belong?\n\"It must be replaced.\"",
              ["[1]", "[2]", "[3]", "[4]"],
              "B", "Sau câu nói mái hiên đã quá cũ và hỏng không thể sửa chữa được ('worn beyond repair'), câu 'It must be replaced' là hoàn toàn phù hợp.")
         ]),

        # 161-164
        ("Questions 161-164 refer to the following advertisement.\n\n"
         "Position: Assistant Editor\nDate Posted: March 15\n"
         "Description: Goldhorse Press, an independent book publishing company focusing on North American gardening, "
         "has provided gardening advice for home gardeners for over 50 years. We are currently seeking an assistant editor to join our expanding team in Charlotte, North Carolina. "
         "We offer an excellent benefits package that includes medical and dental insurance.\n\n"
         "Responsibilities: Assist editors in the acquisition of titles with mainstream appeal; develop manuscripts, collaborate with outside support, and work closely with two managing editors.\n\n"
         "Requirements/Qualifications:\n• Bachelor's degree in a related field\n• Minimum of one year of experience in the publishing industry\n"
         "• Ability to pay close attention to detail\n• Comfortable working as a team member\n• Some experience in gardening preferred\n• Immediate availability\n\n"
         "E-mail résumé and salary requirements to humanresources@goldhorsepress.com.",
         [
             (161, "What title would most likely be published by Goldhorse Press?",
              ["Growing Your Baking Business", "A Tour Guide to North Carolina", "Planting Perennial Flowers", "The Efficient Executive"],
              "C", "Nhà xuất bản chuyên về làm vườn ('gardening') nên tựa sách phù hợp là 'Planting Perennial Flowers'."),
             (162, "What is indicated about Goldhorse Press?",
              ["It is hiring two assistant editors.", "It provides insurance to employees.", "It publishes trade journals.", "It is a newly established company."],
              "B", "Mục đãi ngộ có bảo hiểm y tế và nha khoa: 'includes medical and dental insurance'."),
             (163, "What is a requirement for the advertised position?",
              ["A passion for gardening", "An aptitude for noticing details", "A master's degree in a related field", "An ability to work without supervision"],
              "B", "Yêu cầu có khả năng chú ý tiểu tiết: 'Ability to pay close attention to detail'."),
             (164, "What are applicants asked to do?",
              ["Submit a job application form", "Submit a list of references", "Indicate availability", "Indicate desired pay"],
              "D", "Ứng viên được yêu cầu gửi 'salary requirements' -> mức lương mong muốn.")
         ]),

        # 165-168
        ("Questions 165-168 refer to the following online chat session.\n\n"
         "Derek Marshall [8:19 A.M.]: Hi, everyone. I wanted to give an update on our merger with Ridgewood, Inc., and see if you have had any meetings.\n"
         "Mai Chung [8:20 A.M.]: Have we worked out which of our Derek Boutique locations will stay open?\n"
         "Derek Marshall [8:21 A.M.]: Yes, Derek Boutique will keep 35 stores open with the full line of clothing. The other 12 will move their inventory to the Ridgewood, Inc., locations.\n"
         "Nikita Tamboli [8:22 A.M.]: When should the moves be scheduled? This month?\n"
         "Derek Marshall [8:23 A.M.]: No, the 12 locations don't need to be vacated until the end of next month. Could you schedule this for five weeks from now?\n"
         "Anthony Rossi [8:24 A.M.]: I met over lunch with my managers earlier this week, including the two that I just hired.\n"
         "Nikita Tamboli [8:25 A.M.]: That would work.\n"
         "Mai Chung [8:27 A.M.]: I'll be meeting with the managers on my staff, too, to review the transition plan.\n"
         "Derek Marshall [8:28 A.M.]: This is good work, everyone.\n"
         "Anthony Rossi [8:29 A.M.]: When will the whole process be complete?\n"
         "Derek Marshall [8:30 A.M.]: Within about 6 months.",
         [
             (165, "Where do the writers most likely work?",
              ["At a marketing firm", "At a clothing company", "At a real estate agency", "At a newspaper publisher"],
              "B", "Đoạn chat nhắc tới chuỗi cửa hàng Derek Boutique và 'full line of clothing' -> công ty thời trang/may mặc."),
             (166, "How many stores will be closing?",
              ["5", "6", "12", "35"],
              "C", "35 cửa hàng được giữ lại, 12 cửa hàng còn lại sẽ di dời hàng tồn kho và đóng cửa."),
             (167, "At 8:25 A.M., what does Ms. Tamboli most likely mean when she writes, 'That would work'?",
              ["The current inventory will be doubled.", "Ridgewood, Inc., will close in two weeks.", "Mr. Marshall will meet with the managers.", "She can schedule the moves in the proposed time frame."],
              "D", "Derek đề xuất dời lịch sang 5 tuần nữa và Nikita phản hồi rằng mốc thời gian đó khả thi."),
             (168, "What is suggested about Ms. Chung and Mr. Rossi?",
              ["They supervise other employees.", "They often shop at Ridgewood, Inc.", "They recently went to lunch together.", "They have concerns about the merger."],
              "A", "Cả hai đều nhắc đến việc làm việc/quản lý với các 'managers on my staff' -> họ có vai trò giám sát nhân viên.")
         ]),

        # 169-171
        ("Questions 169-171 refer to the following article.\n\n"
         "From the Shadows to the Limelight - By Calum Ellwood\n"
         "Dr. Esther Nujoma, an agricultural biotechnologist with the Namibia Institute of Applied Sciences, is the author of several books on the practical applications of biotechnology. - [1] -. "
         "According to book critic Paige Kinnock of the London Daily Register, \"Dr. Nujoma has increased the public's awareness of the role of biotechnology in daily life through her ability to translate highly complex scientific material into simple language.\"\n"
         "Her latest work, Shining Behind Shadows, marks a departure from her usual subject matter. - [2] -. Rather, the book highlights the lives and careers of twelve of her peers from Africa and Asia. "
         "The idea came to her three years ago at a conference in Chile. As Dr. Nujoma recalls, \"Listening to one speaker after another, I realized that many of my colleagues are from regions of the world, such as Africa and Asia, that were underrepresented.\" - [3] -.\n"
         "The book is quite compelling. Dr. Nujoma brings to life the stories of how her peers were drawn to the profession and the effort they pour into their work. "
         "The book does, however, fall short in one respect: it provides insights only into the lives and careers of those working in the field of agricultural biotechnology. "
         "I would have welcomed the stories of those specializing in animal, marine, or medical biotechnology, too. - [4] -. Even so, Dr. Nujoma has again succeeded in creating a work that speaks to experts and laypeople alike.",
         [
             (169, "What is NOT indicated about Dr. Nujoma?",
              ["She has visited Chile.", "She is a talented writer.", "She is based out of Namibia.", "She was interviewed by the London Daily Register."],
              "D", "Paige Kinnock thuộc tờ London Daily Register viết bài đánh giá sách chứ bài viết không hề nhắc việc phỏng vấn bà."),
             (170, "What does Mr. Ellwood say about Shining Behind Shadows?",
              ["It focuses on scientists from Africa and Asia.", "It was released three years ago.", "It details why Dr. Nujoma chose her career.", "It describes various branches of biotechnology."],
              "A", "Cuốn sách kể về cuộc đời và sự nghiệp của 12 nhà khoa học đến từ châu Phi và châu Á."),
             (171, "In which of the positions marked [1], [2], [3], and [4] does the following sentence best belong?\n\"Its focus is not on applying biotechnology in real-life situations.\"",
              ["[1]", "[2]", "[3]", "[4]"],
              "B", "Vị trí [2] đứng sau câu nói tác phẩm mới rẽ hướng so với chủ đề quen thuộc, và trước câu 'Rather, the book highlights...'.")
         ]),

        # 172-175
        ("Questions 172-175 refer to the following notice.\n\n"
         "KINGSTON-GARNET ISLAND PASSENGER FERRY SERVICE - General Information\n"
         "Beginning on 15 May, the Kingston-Garnet Island Passenger Ferry Service will resume service for eight weeks during the summer season. "
         "Ferries run daily every half hour from 7 a.m. to 8 p.m. The last ferry to Garnet Island will depart at 7:30 p.m. The last ferry from Garnet Island will leave at 8 p.m.\n"
         "Bicycles are permitted on the passenger ferries. Bicyclists should arrive 30 minutes prior to departure and wait in the special bicycle lane to be loaded first. "
         "There are five racks that hold 50 bicycles on every passenger ferry.\n"
         "No motorized vehicles are permitted on Garnet Island. Overnight parking is allowed in the main ferry terminal lot in Kingston. "
         "Rates are $5 per hour for up to 4 hours and a flat fee of $25 for four to 24 hours.\n"
         "Visit our Web site at www.kgferryservice.com for photographs of the ferry boats, a map of Garnet Island, lists of local attractions on the island, and information about peak-hour fare increases and group discounts.",
         [
             (172, "What is indicated about the ferry service?",
              ["It is available only seasonally.", "Its boats were recently upgraded.", "It takes an hour to reach the island.", "It runs more frequently on weekends."],
              "A", "Phà chỉ chạy trong 8 tuần mùa hè ('eight weeks during the summer season') -> hoạt động theo mùa."),
             (173, "What is true about bicyclists on the ferries?",
              ["They are last to board the boat.", "They must purchase a special ticket.", "They cannot travel on the 7:30 P.M. trip.", "They should arrive at the terminal early."],
              "D", "Người đi xe đạp cần đến trước 30 phút ('arrive 30 minutes prior to departure')."),
             (174, "What is not allowed on Garnet Island?",
              ["Renting bicycles", "Driving cars", "Camping overnight", "Taking photographs"],
              "B", "Hòn đảo cấm các phương tiện gắn động cơ: 'No motorized vehicles are permitted'."),
             (175, "What is indicated about the ferry tickets?",
              ["They can be purchased at stores in Kingston.", "They are less expensive for children.", "They vary in price depending on the time of travel.", "They are more expensive if purchased on the boat."],
              "C", "Giá vé thay đổi theo khung giờ cao điểm: 'information about peak-hour fare increases'.")
         ]),

        # 176-180
        ("Questions 176-180 refer to the following receipt and e-mail.\n\n"
         "Receipt - Green Stripe Press\nOrder Number: GSP20896 | Customer: Jason Ho <jasonho@shoeboxmountain.com>\n"
         "Order Date: December 14 (PREPAID: online order) | Expected Delivery: December 18-20\n"
         "• 1 x CAL201 Complimentary Wild Animals Calendar: $0.00\n"
         "• 7 x ARB132 Accounting Record Book ($19.99 each): $139.93\n"
         "Subtotal: $139.93 | Tax @6%: $8.40 | Shipping & Handling: $0.00 | Total: $148.33\n"
         "Note: There is no charge for shipping and handling for corporate accounts.\n\n"
         "--- E-Mail ---\nTo: customerservice@greenstripepress.com\nFrom: Jason Ho <jasonho@shoeboxmountain.com>\nDate: December 18\nSubject: Order #GSP20896\n\n"
         "Dear Green Stripe Press,\n"
         "I am writing regarding my most recent order (#GSP20896), which was delivered today. If you check my original order, you will see that I ordered six accounting record books. "
         "You sent us seven copies. I'd like to return the one I didn't order and have our corporate credit card refunded, together with the shipping cost the return will incur. Let me know how you would like me to proceed.\n"
         "On a different note, congratulations on the calendar you included in my order! The photos are even more stunning than those in the Ancient Castles calendar you sent us last year. "
         "Some staff members saw mine and want copies of their own. Would you mind sending two more our way?\n\n"
         "Best,\nJason Ho, Owner, Shoebox Mountain",
         [
             (176, "Why was Mr. Ho not charged a shipping fee?",
              ["He took advantage of a promotion.", "He made the purchase for his company.", "He picked up his order in person.", "He overpaid for shipping on a previous order."],
              "B", "Hóa đơn ghi rõ: miễn phí vận chuyển cho tài khoản doanh nghiệp ('corporate accounts')."),
             (177, "What is true about Mr. Ho's order?",
              ["It was damaged in transit.", "It included fragile items.", "It was paid for by check.", "It was delivered on time."],
              "D", "Dự kiến giao ngày 18-20/12, hàng giao đúng ngày 18/12 -> đúng hạn."),
             (178, "According to the e-mail, what was Mr. Ho sent by mistake?",
              ["A castle book he did not order", "A calendar for last year", "An incorrect refund check", "An extra accounting book"],
              "D", "Đặt 6 cuốn sổ kế toán nhưng shop giao thừa 1 cuốn (tổng cộng 7 cuốn)."),
             (179, "What is one reason Mr. Ho wrote the e-mail?",
              ["To offer praise for an item", "To complain about a price", "To order some photo albums", "To recommend a graphic designer"],
              "A", "Khen ngợi bộ lịch tặng kèm ảnh chụp rất đẹp: 'The photos are even more stunning'."),
             (180, "In the e-mail, the phrase 'our way' in paragraph 2, line 4, is closest in meaning to",
              ["in our style", "to our address", "at our expense", "for our benefit"],
              "B", "'sending two more our way' đồng nghĩa với gửi thêm 2 bản tới địa chỉ của chúng tôi.")
         ]),

        # 181-185
        ("Questions 181-185 refer to the following contact form and e-mail.\n\n"
         "Contact Form - Hilgrave Historical Commission\n"
         "Name: Roger Witmondt (roger@witmondtlocations.com)\n"
         "Comment: I work as a movie location scout. At this time, I need to find a location for a client's short film. "
         "The film will be set in the early 1930s, and I understand that Hilgrave has some interesting architecture from that era. "
         "Could you suggest places I should visit when I am in Hilgrave next month? Specifically, I am looking for a vacant building with an old-fashioned storefront "
         "that has distinctive details like stripe-patterned awnings and framed display windows. While having easy access to electricity would be ideal, it is not critical; "
         "my client has a generator that can be used if necessary. My client's aesthetic requirements are the most important consideration. Thank you for your assistance.\n\n"
         "--- E-Mail ---\nTo: roger@witmondtlocations.com\nFrom: brandi_schaertl@hilgravehistoricalcommission.org\nDate: 28 September\nSubject: The information you requested\n\n"
         "Dear Mr. Witmondt,\n"
         "Thank you for reaching out to the Hilgrave Historical Commission. The attached brochure lists the main areas of historical interest in Hilgrave. "
         "Some buildings may not meet all your requirements, but they might be worth considering.\n"
         "One building that is not on the attached list is the old record store at 188 Main Street. It has been closed to business for the past ten years or so, "
         "but it seems to meet your criteria. The owner, Luke Nylund, is currently using it for storage; I would be happy to put you in touch with him. "
         "I doubt you will have trouble getting permission from him to use the space. While Hilgrave was once a bustling town, these days there are fewer businesses in operation. "
         "Many of the town's residents would appreciate the attention a short film might bring to the area.\n\n"
         "Best regards,\nBrandi Schaertl",
         [
             (181, "Why is Mr. Witmondt looking for a filming location in Hilgrave?",
              ["It is known for its scenic mountain views.", "It is a short distance from his office.", "Its business district has appeared in other films.", "Its buildings represent a particular time period."],
              "D", "Tìm kiếm địa điểm có kiến trúc tiêu biểu của đầu thập niên 1930."),
             (182, "On the contact form, the word 'critical' in paragraph 2, line 4, is closest in meaning to",
              ["judgmental", "essential", "sustainable", "available"],
              "B", "'not critical' mang nghĩa không phải là điều kiện thiết yếu / cốt tử ('not essential')."),
             (183, "What is implied about 188 Main Street?",
              ["It was once used as a residence.", "It is frequently visited by tourists.", "It has several floors.", "It has decorative design features."],
              "D", "Brandi bảo tòa nhà này thỏa mãn tiêu chí của Roger (mặt tiền phong cách cổ điển, mái che có họa tiết...)."),
             (184, "What does Ms. Schaertl offer to do?",
              ["Advertise a forthcoming film", "Arrange to have a building cleaned", "Connect Mr. Witmondt with a building's owner", "Help Mr. Witmondt acquire necessary permits from the town"],
              "C", "Cô ấy sẵn lòng kết nối Roger với chủ sở hữu Luke Nylund: 'put you in touch with him'."),
             (185, "What does Ms. Schaertl suggest that Hilgrave needs?",
              ["More publicity", "Additional parking", "A storage facility", "A business directory"],
              "A", "Thị trấn cần được quảng bá rộng rãi hơn ('appreciate the attention a short film might bring').")
         ]),

        # 186-190
        ("Questions 186-190 refer to the following e-mails and Web page.\n\n"
         "--- E-Mail 1 ---\nTo: Cornelia Payne <cpayne@roughwing.co.uk>\nFrom: Pragya Mehta <pmehta@airsky.in>\nDate: 18 January\nSubject: Lecture\n\n"
         "Dear Cornelia,\nI am sorry for the late notice, but I will not be able to join you for the lecture. My new position at the university in Mumbai requires me to remain on campus. "
         "Although I know you are fully prepared to deliver the lecture on your own, I have attached a copy of the notes I had prepared for the presentation. "
         "After you review them, let me know if there is anything else I can add. I had very much looked forward to traveling to Freeport, The Bahamas, for the first time "
         "and to seeing you again. I certainly miss working with you at the City University. Good luck with the lecture.\nPragya\n\n"
         "--- E-Mail 2 ---\nTo: Ezra Halton <ehalton@ansonhouse.org>\nFrom: Cornelia Payne <cpayne@roughwing.co.uk>\nDate: 20 January\nSubject: Information\n\n"
         "Dear Mr. Halton,\nI am pleased to confirm that 19 March works well for the lecture at the Anson House. As I had mentioned to you over the phone, I will be giving the lecture on my own. "
         "My publisher, Alphagamma Press, will be sending 30 copies of my most recent work, coauthored with Dr. Pragya Mehta, to your institution. "
         "You and I will both receive an e-mail confirmation from Alphagamma when the books ship. They should arrive at least a week before the lecture. "
         "I look forward to meeting you soon.\nCornelia Payne\n\n"
         "--- Web Page ---\nSpeaker Series - The Anson House (Freeport, The Bahamas)\n"
         "Tickets: $35 each or $90 for the series of three. Complimentary refreshments served. Reservations are not required but are recommended.\n"
         "• 7 March, 7:00-8:30 P.M.: Ms. Janelle Pears\n• 12 March, 7:00-8:30 P.M.: Mr. Gregory Li\n"
         "• 19 March, 7:00-8:30 P.M.: Dr. Cornelia Payne (Dr. Payne discusses her most recently published book, The Transatlantic World of the Nineteenth Century, "
         "coauthored with Dr. Pragya Mehta. Dr. Payne has taught at the City University of Stoke-on-Trent in England for nearly 25 years.)",
         [
             (186, "Why did Dr. Mehta send the e-mail to Dr. Payne?",
              ["To cancel a vacation plan", "To request lecture notes", "To confirm a meeting", "To offer an apology"],
              "D", "Xin lỗi vì thông báo muộn không thể đến tham gia buổi thuyết giảng được ('I am sorry for the late notice')."),
             (187, "According to the second e-mail, what did Dr. Payne do?",
              ["Arrange a delivery", "Give Mr. Halton a gift", "Revise her presentation", "Announce an address change"],
              "A", "Sắp xếp nhà xuất bản gửi 30 cuốn sách tới viện bảo tàng ('sending 30 copies... to your institution')."),
             (188, "What is indicated about Alphagamma Press?",
              ["It has headquarters in The Bahamas.", "It regularly ships materials to Anson House.", "It published The Transatlantic World of the Nineteenth Century.", "It is paying Dr. Payne's travel and accommodation expenses."],
              "C", "Là đơn vị xuất bản cuốn sách mới nhất của Dr. Payne và Dr. Mehta."),
             (189, "Where were Dr. Mehta and Dr. Payne most likely colleagues?",
              ["In Mumbai", "In Freeport", "In Stoke-on-Trent", "In Charleston"],
              "C", "Email 1 nhắc 'miss working with you at the City University', trang web nêu rõ đó là City University of Stoke-on-Trent."),
             (190, "What is indicated about the Speaker Series in the Web page?",
              ["Discounts are unavailable.", "Reservations are optional.", "Events are held in the morning.", "Refreshments are not included."],
              "B", "Trang web ghi: 'Reservations are not required but are recommended' -> việc đặt chỗ là tùy chọn (optional).")
         ]),

        # 191-195
        ("Questions 191-195 refer to the following advertisement, e-mail, and press release.\n\n"
         "--- Advertisement ---\nCRYN Group: We find the best employees for your company.\nPosted: November 25\n"
         "Position titles: Director of Operations (TL0015), Marketing Director (TL0023), Quality Control Director (TL0027), Director of Category Management (TL0045)\n"
         "To apply: Send your résumé to mdoro@cryngroup.ca.com with posting number 2098 in the subject line.\n"
         "Company: Our client plans to chart an aggressive growth path in Latin America, where it plans to start operations next year.\n"
         "Qualifications: Candidates must have a formal business degree and a proven management record in at least one international setting. Experience with online sales and marketing preferred.\n\n"
         "--- E-Mail ---\nTo: Sven Arvidson <sarvidson@barkent.de.com>\nFrom: Maria Doro <mdoro@cryngroup.ca.com>\nDate: December 10\nRe: Posting number 2098\n\n"
         "Dear Mr. Arvidson,\nThank you for submitting your résumé. I'd like to schedule a preliminary telephone conversation with you as soon as possible "
         "to determine whether you would be a good choice for our client. I would particularly like to discuss whether you would be prepared to work at our client's new facility overseas. "
         "Let me know if you are available for a 30-minute phone call sometime between 10:00 A.M. and 2:00 P.M. EST on Monday or Tuesday of next week.\nMaria Doro\n\n"
         "--- Press Release ---\nTayerson Ltd. Names New Director\n"
         "Toronto, April 5 - Tayerson Ltd. continues to staff its new international outpost, which opens one month from today. In this case, the lucky candidate is Sven Arvidson. "
         "\"Category Management is a relatively new area and one that I'm sure I'll enjoy exploring and developing as director,\" Mr. Arvidson said. "
         "\"As a major online marketer of nutritional supplements, Tayerson is poised to lead the way toward better living...\"",
         [
             (191, "According to the advertisement, what qualification is not required of applicants?",
              ["A degree in business", "Previous employment in a managerial position", "Experience working overseas", "Online sales and marketing experience"],
              "D", "Kinh nghiệm bán hàng online chỉ là 'preferred' (ưu tiên) chứ không phải bắt buộc ('must have')."),
             (192, "What most likely is the location of the position for which Mr. Arvidson applied?",
              ["Latin America", "North America", "Europe", "Asia"],
              "A", "Thông báo tuyển dụng ghi cơ sở mới chuẩn bị mở vào năm tới tại Latin America."),
             (193, "What does Ms. Doro ask Mr. Arvidson to provide in his reply?",
              ["Information on his leadership skills", "His availability for an interview", "A current résumé", "The names of two references"],
              "B", "Hỏi xem anh ấy có rảnh vào khung giờ nào để gọi điện phỏng vấn sơ bộ: 'availability for a 30-minute phone call'."),
             (194, "What job code did Mr. Arvidson most likely reference in his application?",
              ["TL0015", "TL0023", "TL0027", "TL0045"],
              "D", "Sven Arvidson nhận vị trí Category Management, tương ứng với mã job TL0045."),
             (195, "According to the press release, what does Tayerson Ltd. sell?",
              ["Medical devices", "Accounting software", "Nutritional supplements", "Exercise equipment"],
              "C", "Bài báo ghi: 'major online marketer of nutritional supplements' -> thực phẩm bổ sung dinh dưỡng.")
         ]),

        # 196-200
        ("Questions 196-200 refer to the following brochure, e-mail, and schedule.\n\n"
         "--- Brochure ---\nFGJ's Business Expert Series: Prospect Data Purchasing (45-Minute Webinar, 11 May, 2:30 P.M.)\n"
         "Marketing professionals in particular rely on accurate data about potential customers to be successful. In this Webinar, "
         "Briana Carrera, chief financial officer of Pile One Market Data, explains how to avoid the pitfalls of acquiring incomplete data...\n\n"
         "--- E-Mail ---\nTo: Sandra Lescure\nFrom: Gino Stelleti\nDate: 12 May\nSubject: FGJ's Webinar\n\n"
         "Dear Sandra,\nYou had asked me to give you a report on the Webinar I attended yesterday about acquiring data. To be honest, "
         "even though the facilitator was clearly knowledgeable, she did not tell me anything I didn't already know. I was hoping that she would show us "
         "the differences between the many types of databases out there, but she never touched on that topic. I am not so sure all these Webinars are useful; "
         "in the end, they always seem to be steering us to purchase from one particular company or another in this case it was Pile One. "
         "Having said that, the new schedule just came out, and there is another Webinar coming up soon that I am hoping will provide the information I want "
         "about how to store and sort data. I have already signed up for it; I will keep you posted.\nGino\n\n"
         "--- Schedule ---\nFGJ's Upcoming 45-Minute Webinars:\n"
         "• Basics of Market Research | 1 June, 9:30 A.M. | Ed Quinones | €12.00\n"
         "• Tips for Accelerating Sales | 12 June, 9:30 A.M. | Cameron Stone | €16.00\n"
         "• What is Market Automation? | 5 July, 11:00 A.M. | Ed Quinones | €12.00\n"
         "• Choosing the Right Database | 17 July, 2:00 P.M. | Selina Tucci | €12.00",
         [
             (196, "According to the brochure, who would benefit most from the Prospect Data Purchasing Webinar?",
              ["A financial analyst", "A marketing manager", "A customer-service representative", "An information-technology specialist"],
              "B", "Tài liệu quảng cáo ghi: 'Marketing professionals in particular rely on accurate data'."),
             (197, "What was Mr. Stelleti's complaint about the Webinar?",
              ["He did not like the format.", "He did not learn anything new.", "He had difficulty hearing everything.", "He found the topic to be too complex."],
              "B", "Anh ấy phàn nàn: 'she did not tell me anything I didn't already know' -> không học thêm được điều gì mới."),
             (198, "What does Mr. Stelleti suggest about Ms. Carrera?",
              ["She recently joined a new company.", "She has experience managing databases.", "She had asked him to sign up for the Webinar.", "She tried to sell her company's services."],
              "D", "Gino nhận xét buổi hội thảo hướng người nghe mua dịch vụ từ công ty của diễn giả: 'steering us to purchase from... Pile One'."),
             (199, "When will Mr. Stelleti most likely attend another Webinar?",
              ["On June 1", "On June 12", "On July 5", "On July 17"],
              "D", "Gino muốn tìm hiểu về cách lưu trữ và phân loại dữ liệu ('store and sort data'), tương ứng với chủ đề 'Choosing the Right Database' vào ngày 17/7."),
             (200, "According to the schedule, what do all of the Webinars have in common?",
              ["They cost the same amount.", "They are held in the morning.", "They last the same amount of time.", "They are recorded for future playback."],
              "C", "Tất cả các hội thảo đều kéo dài trong cùng một khoảng thời gian: 45 phút ('Upcoming 45-Minute Webinars').")
         ])
    ]

    for passage_text, questions in part7_sections:
        doc.add_paragraph("[PASSAGE_START]")
        doc.add_paragraph(passage_text)
        doc.add_paragraph("[PASSAGE_END]")

        for q_num, text, opts, key, hint in questions:
            doc.add_paragraph(f"[Q: {q_num}] {text}")
            doc.add_paragraph("[SHUFFLE_Q: FALSE]")
            doc.add_paragraph("[SHUFFLE_A: TRUE]")
            for idx, opt_label in enumerate(["A", "B", "C", "D"]):
                doc.add_paragraph(f"[{opt_label}] {opts[idx]}")
            doc.add_paragraph(f"[KEY: {key}]")
            doc.add_paragraph(f"[HINT: {hint}]")
            doc.add_paragraph("[END_Q]")
            doc.add_paragraph("")

    # Lưu file Word
    file_name = "ETS_Test_09_Tagged.docx"
    doc.save(file_name)
    print(f"--> Đã tạo thành công file Word đề thi chuẩn Tags: {file_name}")

if __name__ == "__main__":
    create_toeic_reading_test_09()